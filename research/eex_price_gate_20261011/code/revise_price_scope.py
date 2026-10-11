"""Make a separate v8.1 literature/scope revision, preserving every existing equation/table/image."""
from pathlib import Path
import argparse,copy,hashlib,json,zipfile
from lxml import etree as ET
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
SOURCE_SHA='15180678628a21a511b6f2ff49c91b1eb5bd89168129b522585ad20cc2452ee6'
REPLACEMENTS=[
 ('Bosco (2023) studies Phase III',
  'Bosco (2023) studies Phase III EUA auction prices and returns using GARCH specifications. It includes participation, demand coverage and the difference between maximum and minimum submitted prices. This bid-price spread is not a quantity-dispersion measure. Niu and Liu (2024) instead forecast EUA futures volatility. Our outcomes are quantity statistics from completed auctions; no price-premium equation is estimated.'),
 ('External price informativeness is untested.',
  'Price informativeness remains untested. Assessing an incremental association beyond participation, demand coverage and bid-price spread requires a compatible secondary-market quote at a separately documented bidding-window close. The previously inspected report supplies monthly rather than usable per-auction reference comparisons (European Commission 2023). Primary-auction prices have now been extracted, but the workbook field Time is not an authenticated deadline. No eligible external series was acquired and no price-premium regression was run. This is a data limitation, not a null result. The exact-2026-snapshot limitation is stated below.')]
NEW_PARAGRAPH='Otsubo (2026) examines EUA futures, a carbon ETF and auction prices, aligning futures and ETF prices to 16:00 UK time in its daily analysis. That cross-venue comparison is distinct from matching each auction to a quote immediately before bidding closes; its daily alignment cannot supply the required synchronous reference.'
REFERENCE='Otsubo Y (2026) Intraday and daily price discovery in carbon markets: EUA futures versus carbon ETFs. International Review of Finance 26:e70099. https://doi.org/10.1111/irfi.70099'

def math_records(p):
 with zipfile.ZipFile(p) as z:
  root=ET.fromstring(z.read('word/document.xml'))
  ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
  math=[ET.tostring(x,method='c14n') for x in root.xpath('//m:oMath',namespaces=ns)]
  tables=[ET.tostring(x,method='c14n') for x in root.xpath('//w:tbl',namespaces=ns)]
  media={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
  return math,tables,media

def plain_replace(p,text):
 if p._p.xpath('.//m:oMath'):raise ValueError('Refusing to overwrite paragraph with math')
 template=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();run=p.add_run(text)
 if template is not None:run._r.insert(0,template)

def revise(source:Path,out:Path):
 if hashlib.sha256(source.read_bytes()).hexdigest()!=SOURCE_SHA:raise ValueError('Unexpected v8 source')
 out.mkdir(parents=True,exist_ok=True);doc=Document(source);changes=[]
 for prefix,text in REPLACEMENTS:
  found=[p for p in doc.paragraphs if p.text.startswith(prefix)]
  if len(found)!=1:raise ValueError('Ambiguous target '+prefix)
  p=found[0];changes.append({'old':p.text,'new':text});plain_replace(p,text)
  if prefix.startswith('Bosco'):
   element=copy.deepcopy(p._p);p._p.addnext(element);np=Paragraph(element,p._parent);plain_replace(np,NEW_PARAGRAPH)
 found=[p for p in doc.paragraphs if p.text.startswith('Niu H, Liu T (2024)')]
 if len(found)!=1:raise ValueError('Reference anchor missing')
 anchor=found[0];element=copy.deepcopy(anchor._p);anchor._p.addnext(element);plain_replace(Paragraph(element,anchor._parent),REFERENCE)
 for section in doc.sections:
  for p in section.header.paragraphs:
   for run in p.runs:run.text=run.text.replace('AUTHOR-REVIEW V8','AUTHOR-REVIEW V8.1')
 doc.core_properties.subject='v8.1 literature and price-scope revision; no price estimates or changes to existing results'
 doc.core_properties.comments='No funding confirmed; prior Wang review retained. E2 reference gate remains unsatisfied.'
 dest=out/'MANUSCRIPT_EN_v8_1.docx';doc.save(dest)
 a=math_records(source);b=math_records(dest)
 if a!=b:raise ValueError('Protected math, tables or embedded images changed')
 s={'date':'2026-10-11','source_commit':'b23920a1ed51bdd79428e4ebef5dd1bd2b9607e4','source_sha256':SOURCE_SHA,
 'output_sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'native_math_objects_preserved':len(a[0]),
 'physical_table_objects_preserved':len(a[1]),'embedded_images_preserved':len(a[2]),
 'existing_mathematics_tables_images_identical':True,'new_price_regressions':0,'v8_overwritten':False,
 'changes':changes,'added_paragraph':NEW_PARAGRAPH,'added_reference':REFERENCE}
 (out/'MANUSCRIPT_CHANGELOG.json').write_text(json.dumps(s,ensure_ascii=False,indent=2))
 return s
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
 print(json.dumps(revise(a.source,a.out),indent=2,ensure_ascii=False))
