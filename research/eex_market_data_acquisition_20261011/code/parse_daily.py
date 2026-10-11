"""Read author-provided daily date/Y rows; this does not authenticate prices."""
import io
import csv
import stat
from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook


def rows_from_bytes(body):
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError("The accepted author file must contain exactly one sheet")
        worksheet = workbook[workbook.sheetnames[0]]
        worksheet.reset_dimensions()
        source_rows = list(worksheet.iter_rows(values_only=True))
        if not source_rows or list(source_rows[0]) != ["date", "Y"]:
            raise ValueError("Literal original headers differ from date/Y")
        rows = []
        for source_row, cells in enumerate(source_rows[1:], start=2):
            observed, reported = cells
            if isinstance(observed, datetime):
                observed = observed.date()
            if not isinstance(observed, date):
                raise ValueError("Unexpected original date cell")
            rows.append({"source_sheet": worksheet.title, "source_row": source_row,
                         "date": observed.isoformat(), "reported_Y": reported})
        return rows
    finally:
        workbook.close()


def csv_bytes(rows):
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=["source_sheet", "source_row", "date", "reported_Y"], lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8")


def audit(rows):
    dates = [row["date"] for row in rows]
    if not dates:
        raise ValueError("No original date observations")
    prices = [row["reported_Y"] for row in rows]
    return {"observation_count": len(rows), "date_min": min(dates), "date_max": max(dates),
            "unique_date_count": len(set(dates)), "duplicate_date_count": len(dates) - len(set(dates)),
            "dates_strictly_increasing": all(a < b for a, b in zip(dates, dates[1:])),
            "missing_price_count": sum(value is None for value in prices),
            "zero_price_count": sum(value == 0 for value in prices if value is not None),
            "year_row_counts": dict(sorted(Counter(value[:4] for value in dates).items())),
            "unit": None, "specific_allowance_and_contract": None, "venue": None,
            "timestamp_timezone": None, "price_measurement": "daily_close_claimed_by_author_metadata_not_independently_authenticated",
            "source_grade": "author_repository_copy", "auction_close_eligible_reference_rows": 0}


if __name__ == "__main__":
    import argparse
    import hashlib
    import json
    from pathlib import Path

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--write", action="store_true", help="Regenerate only the derived CSV and coverage audit")
    arguments = parser.parse_args()
    root = arguments.root.resolve()
    contract = json.loads((root / "SOURCE_CONTRACT.json").read_text(encoding="utf-8"))
    inputs = contract.get("sources")
    if not isinstance(inputs, list) or len(inputs) != 1:
        raise ValueError("Exactly one frozen original input is required")
    record = inputs[0]
    source = (root / record["path"]).resolve()
    if not source.is_relative_to(root):
        raise ValueError("Original input escapes the package")
    body = source.read_bytes()
    if (len(body) != record["bytes"] or hashlib.sha256(body).hexdigest() != record["sha256"]
            or hashlib.md5(body).hexdigest() != record["publisher_md5"]):
        raise ValueError("Original input differs from the frozen source contract")
    rows = rows_from_bytes(body)
    output = csv_bytes(rows)
    result = audit(rows)
    result["original_sha256"] = record["sha256"]
    result["original_headers"] = ["date", "Y"]
    result["original_sheet"] = rows[0]["source_sheet"]
    csv_path = root / "data/parsed/figshare_eu_daily.csv"
    audit_path = root / "results/DAILY_AUDIT.json"
    if arguments.write:
        audit_output = (json.dumps(result, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        destinations = ((csv_path, output), (audit_path, audit_output))
        existing = set()
        for destination, expected in destinations:
            if not destination.resolve().is_relative_to(root):
                raise ValueError("Derived output escapes the package")
            component = root
            for part in ("", *destination.relative_to(root).parts):
                if part:
                    component = component / part
                try:
                    status = component.lstat()
                except FileNotFoundError:
                    continue
                if stat.S_ISLNK(status.st_mode) or getattr(status, "st_file_attributes", 0) & 0x400:
                    raise ValueError("Derived output contains a link/reparse component")
            if destination.exists():
                if destination.samefile(source):
                    raise ValueError("Derived output aliases the frozen original")
                if not destination.is_file() or destination.stat().st_nlink != 1:
                    raise ValueError("Derived output must be an unaliased regular file")
                if destination.read_bytes() != expected:
                    raise ValueError("Existing derived output differs; refusing to overwrite")
                existing.add(destination)
        # Both outputs must pass preflight before creating any directory or file.
        for destination, expected in destinations:
            if destination not in existing:
                destination.parent.mkdir(parents=True, exist_ok=True)
                with destination.open("xb") as handle:
                    handle.write(expected)
    else:
        if csv_path.read_bytes() != output:
            raise ValueError("Parsed CSV differs from a fresh original-byte extraction")
        if json.loads(audit_path.read_text(encoding="utf-8")) != result:
            raise ValueError("Coverage audit differs from original rows")
    print(json.dumps({"passed": True, "mode": "write" if arguments.write else "check", "audit": result}, ensure_ascii=False))
