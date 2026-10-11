"""Logic fixtures are synthetic; one CLI check uses the actual public original."""
import importlib.util
import csv
import io
import unittest
import zipfile
import subprocess
import sys
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook


class DailyParserTests(unittest.TestCase):
    def load_parser(self):
        path = Path(__file__).with_name("parse_daily.py")
        if not path.exists():
            self.fail("The original-workbook parser is not implemented")
        spec = importlib.util.spec_from_file_location("daily_parser", path)
        if spec is None or spec.loader is None:
            raise RuntimeError("Cannot load parser")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def test_preserves_zero_and_source_rows(self):
        workbook = Workbook()
        worksheet = workbook.active
        if worksheet is None:
            raise RuntimeError("Fixture worksheet is absent")
        worksheet.title = "logic fixture, not market data"
        worksheet.append(["date", "Y"])
        worksheet.append([datetime(2000, 1, 3), 0])
        worksheet.append([datetime(2000, 1, 4), 12.5])
        buffer = io.BytesIO()
        workbook.save(buffer)
        rows = self.load_parser().rows_from_bytes(buffer.getvalue())
        self.assertEqual([r["source_row"] for r in rows], [2, 3])
        self.assertEqual([r["date"] for r in rows], ["2000-01-03", "2000-01-04"])
        self.assertEqual([r["reported_Y"] for r in rows], [0, 12.5])

    def test_reads_physical_rows_when_declared_dimension_is_A1(self):
        workbook = Workbook()
        worksheet = workbook.active
        if worksheet is None:
            raise RuntimeError("Fixture worksheet is absent")
        worksheet.append(["date", "Y"])
        worksheet.append([datetime(2000, 1, 3), 1])
        worksheet.append([datetime(2000, 1, 4), 2])
        original = io.BytesIO()
        workbook.save(original)
        altered = io.BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(altered, "w") as target:
            for member in source.infolist():
                body = source.read(member)
                if member.filename == "xl/worksheets/sheet1.xml":
                    body = body.replace(b'ref="A1:B3"', b'ref="A1"')
                target.writestr(member, body)
        try:
            rows = self.load_parser().rows_from_bytes(altered.getvalue())
        except ValueError as error:
            self.fail("Stale worksheet dimensions hid real cells: " + str(error))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[-1]["source_row"], 3)

    def test_csv_distinguishes_missing_price_from_numeric_zero(self):
        module = self.load_parser()
        if not hasattr(module, "csv_bytes"):
            self.fail("The lossless CSV serializer is not implemented")
        fixture_rows = [
            {"source_sheet": "fixture", "source_row": 2, "date": "2000-01-03", "reported_Y": None},
            {"source_sheet": "fixture", "source_row": 3, "date": "2000-01-04", "reported_Y": 0},
        ]
        parsed = list(csv.DictReader(io.StringIO(module.csv_bytes(fixture_rows).decode("utf-8"))))
        self.assertEqual([row["reported_Y"] for row in parsed], ["", "0"])
        self.assertEqual([row["source_row"] for row in parsed], ["2", "3"])

    def test_audit_counts_missing_and_zero_separately(self):
        module = self.load_parser()
        if not hasattr(module, "audit"):
            self.fail("The observed-row coverage audit is not implemented")
        fixture_rows = [
            {"source_sheet": "fixture", "source_row": 2, "date": "2000-01-03", "reported_Y": None},
            {"source_sheet": "fixture", "source_row": 3, "date": "2000-01-04", "reported_Y": 0},
        ]
        result = module.audit(fixture_rows)
        self.assertEqual(result["observation_count"], 2)
        self.assertEqual(result["missing_price_count"], 1)
        self.assertEqual(result["zero_price_count"], 1)
        self.assertEqual(result["date_min"], "2000-01-03")
        self.assertEqual(result["date_max"], "2000-01-04")
        self.assertEqual(result["year_row_counts"], {"2000": 2})

    def test_download_cli_verifies_packaged_original_without_network(self):
        # Integration check against the real public original, not a generated market fixture.
        script = Path(__file__).with_name("download_public.py")
        if not script.exists():
            self.fail("The verified download CLI is not implemented")
        result = subprocess.run([sys.executable, "-B", str(script), "--verify-only"], capture_output=True, text=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)
        receipt = json.loads(result.stdout)
        self.assertEqual(receipt["sha256"], "8617b185007c3f27c2bad593636aac2a6afc602d90ab048f5395bde048556985")
        self.assertTrue(receipt["reused_existing_original"])


class DerivedWriteSafetyTests(unittest.TestCase):
    """Temporary logic fixtures only; never write to the real package."""

    def setUp(self):
        import hashlib
        import os
        import tempfile
        self.temporary = tempfile.TemporaryDirectory(prefix="daily_write_logic_", dir=Path(os.environ["TMPDIR"]))
        self.addCleanup(self.temporary.cleanup)
        # Resolve first: Windows temporary paths can contain 8.3 aliases.
        self.root = Path(self.temporary.name).resolve() / "package"
        self.root.mkdir()
        self.source = self.root / "original.xlsx"
        workbook = Workbook()
        workbook.active.title = "logic fixture, not market data"
        workbook.active.append(["date", "Y"])
        workbook.active.append([datetime(2000, 1, 3), 1])
        workbook.save(self.source)
        self.original = self.source.read_bytes()
        contract = {"sources": [{"path": "original.xlsx", "bytes": len(self.original),
                    "sha256": hashlib.sha256(self.original).hexdigest(),
                    "publisher_md5": hashlib.md5(self.original).hexdigest()}]}
        (self.root / "SOURCE_CONTRACT.json").write_text(json.dumps(contract), encoding="utf-8")
        self.csv = self.root / "data/parsed/figshare_eu_daily.csv"
        self.audit = self.root / "results/DAILY_AUDIT.json"

    def run_write(self):
        command = [sys.executable]
        if sys.flags.optimize:
            command.append("-O")
        return subprocess.run(command + ["-B", str(Path(__file__).with_name("parse_daily.py")),
                              "--root", str(self.root), "--write"],
                              capture_output=True, text=True, timeout=15)

    def run_in_process(self):
        import runpy
        from unittest.mock import patch
        script = Path(__file__).with_name("parse_daily.py")
        with patch.object(sys, "argv", [str(script), "--root", str(self.root), "--write"]):
            runpy.run_path(str(script), run_name="__main__")

    def test_resolved_output_escape_rejected_before_creation(self):
        from unittest.mock import patch
        real_resolve = Path.resolve
        outside = self.root.parent / "outside.csv"
        def resolve(path, *args, **kwargs):
            if path == self.csv:
                return outside
            return real_resolve(path, *args, **kwargs)
        with patch.object(Path, "resolve", resolve):
            with self.assertRaisesRegex(ValueError, "escapes"):
                self.run_in_process()
        self.assertFalse(self.csv.parent.exists())
        self.assertFalse(self.audit.parent.exists())
        self.assertFalse(outside.exists())

    def test_reparse_component_rejected_before_creation(self):
        from types import SimpleNamespace
        from unittest.mock import patch
        self.audit.parent.mkdir()
        real_lstat = Path.lstat
        def lstat(path, *args, **kwargs):
            status = real_lstat(path, *args, **kwargs)
            if path == self.audit.parent:
                return SimpleNamespace(st_mode=status.st_mode, st_file_attributes=0x400)
            return status
        with patch.object(Path, "lstat", lstat):
            with self.assertRaisesRegex(ValueError, "link|reparse"):
                self.run_in_process()
        self.assertFalse(self.csv.parent.exists())
        self.assertFalse(self.audit.exists())
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_posix_symlink_component_rejected_before_creation(self):
        import stat
        from types import SimpleNamespace
        from unittest.mock import patch
        self.audit.parent.mkdir()
        real_lstat = Path.lstat
        def lstat(path, *args, **kwargs):
            status = real_lstat(path, *args, **kwargs)
            if path == self.audit.parent:
                return SimpleNamespace(st_mode=stat.S_IFLNK, st_file_attributes=0)
            return status
        with patch.object(Path, "lstat", lstat):
            with self.assertRaisesRegex(ValueError, "link|reparse"):
                self.run_in_process()
        self.assertFalse(self.csv.parent.exists())
        self.assertFalse(self.audit.exists())
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_second_output_conflict_preflight_before_first_creation(self):
        self.audit.parent.mkdir()
        self.audit.write_bytes(b"logic fixture: existing unrelated file")
        result = self.run_write()
        self.assertNotEqual(result.returncode, 0, "Conflicting audit must be rejected before any output write")
        self.assertFalse(self.csv.parent.exists())
        self.assertEqual(self.audit.read_bytes(), b"logic fixture: existing unrelated file")
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_original_hardlink_rejected_without_modification(self):
        import os
        self.csv.parent.mkdir(parents=True)
        os.link(self.source, self.csv)
        result = self.run_write()
        self.assertNotEqual(result.returncode, 0, "Original hardlink must be rejected")
        self.assertEqual(self.source.read_bytes(), self.original)
        self.assertFalse(self.audit.parent.exists())


if __name__ == "__main__":
    unittest.main()
