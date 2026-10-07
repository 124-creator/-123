"""Independent quote-definition checks. No model training or old metrics."""
from decimal import Decimal
import argparse
import csv
import datetime as dt
import hashlib
import json
import html
import re
import zipfile
from pathlib import Path
import openpyxl
import xlrd

def reported_decimal(value):
    if value is None or str(value).strip() in ('','--','—'):
        return None
    return Decimal(str(value).strip().replace(',',''))

def reconstructed_price(date, prices):
    if date >= '2026-01-05':
        return prices['CEA25']
    if date >= '2025-04-29':
        keys = ['CEA','CEA21','CEA22','CEA23','CEA24']
    elif date >= '2024-10-28':
        keys = ['CEA','CEA21','CEA22','CEA23']
    elif date >= '2023-08-28':
        keys = ['CEA','CEA21','CEA22']
    elif date >= '2021-07-16':
        keys = ['CEA']
    else:
        raise ValueError('Date predates the verified national market history')
    return sum((prices[key] for key in keys), Decimal(0)) / Decimal(len(keys))

def export_dataset(root, name, rows, source, extra):
    rows = sorted(rows,key=lambda row:row['date'])
    dates = [row['date'] for row in rows]
    if len(dates)!=len(set(dates)) or not dates:
        raise ValueError('Dataset dates must be nonempty and unique')
    if max(dates)>'2026-10-07':
        raise ValueError('The supplied snapshot contains future observations')
    path = root/'data/parsed'/(name+'.csv')
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        writer = csv.DictWriter(f,fieldnames=list(rows[0]))
        writer.writeheader();writer.writerows(rows)
    return {'dataset':name,'path':path.relative_to(root).as_posix(),'rows':len(rows),'date_min':min(dates),'date_max':max(dates),'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'source_original':source,**extra}

def check_local_exchange_snapshots(root):
    meta=json.loads((root/'metadata/local_exchange_sources.json').read_text(encoding='utf-8'))
    base=root/'data/local_only'
    for record in meta['sources']:
        body=(base/record['filename']).read_bytes()
        if len(body)!=record['bytes'] or hashlib.sha256(body).hexdigest()!=record['sha256']:
            raise ValueError(f"Checksum or length mismatch: {record['filename']}")
    codes={'composite':'cea_composite_live.json','CEA':'cea_19_20_live.json',**{f'CEA{i}':f'cea_vintage_CEA{i}.json' for i in range(21,26)}}
    series={};summary=[]
    for code,file in codes.items():
        rows=json.loads((base/file).read_text(encoding='utf-8'))
        if not all(len(row)==6 for row in rows):
            raise ValueError(f'Unexpected row width: {file}')
        dates=[row[0] for row in rows]
        if len(dates)!=len(set(dates)):
            raise ValueError(f'Duplicate dates: {file}')
        series[code]={row[0]:Decimal(row[2]) for row in rows}
        summary.append({'dataset':code,'rows':len(rows),'date_min':min(dates),'date_max':max(dates),'reported_zero_quantity_rows':sum(Decimal(row[5])==0 for row in rows),'status_certified_as_no_trade':False})
    starts=['2021-07-16','2023-08-28','2024-10-28','2025-04-29','2026-01-05']
    buckets={date:{'effective_from':date,'days':0,'mismatch_days':0,'max_absolute_difference':Decimal(0)} for date in starts}
    missing=[];current=[]
    for date,reported in sorted(series['composite'].items()):
        component_prices={code:data[date] for code,data in series.items() if code!='composite' and date in data}
        try:expected=reconstructed_price(date,component_prices)
        except KeyError:missing.append(date);continue
        bucket=buckets[max(x for x in starts if x<=date)];difference=abs(reported-expected)
        bucket['days']+=1;bucket['mismatch_days']+=int(difference>Decimal('0.01'))
        bucket['max_absolute_difference']=max(bucket['max_absolute_difference'],difference)
        if date>='2026-01-05':current.append(reported==component_prices['CEA25'])
    for bucket in buckets.values():bucket['max_absolute_difference']=str(bucket['max_absolute_difference'])
    all_dates=set().union(*(set(data) for code,data in series.items() if code!='composite'))
    cea={'series':summary,'composite_observations':len(series['composite']),'regimes':list(buckets.values()),'unavailable_component_dates':missing,'composite_missing_dates_with_vintage_observations':sorted(all_dates-set(series['composite'])),'cea25_period_days':len(current),'cea25_period_exact_matches':sum(current),'comparison_tolerance_cny_per_tonne':'0.01','raw_daily_quotes_republished':False,'quantity_zero_does_not_by_itself_establish_no_trade':True}
    gdea=[];totals=[]
    for n in [1,2,3]:
        response=json.loads((base/f'gdea_api_page{n}.json').read_text(encoding='utf-8'))['data']
        totals.append(int(response['total']));gdea.extend(response['rows'])
    if len(set(totals))!=1 or len(gdea)!=totals[0]:
        raise ValueError('GDEA pagination totals do not match the collected rows')
    if len({x['currentDay'] for x in gdea})!=len(gdea):
        raise ValueError('Duplicate GDEA dates')
    if not all(x['productCode']=='000001' and x['productName']=='GDEA' for x in gdea):
        raise ValueError('Unexpected GDEA product identity')
    known_quantity=[x for x in gdea if reported_decimal(x.get('exchangeRate')) is not None]
    missing_quantity=[x['currentDay'] for x in gdea if reported_decimal(x.get('exchangeRate')) is None]
    conflicts=[x['currentDay'] for x in known_quantity if x.get('message')=='当日无成交' and reported_decimal(x['exchangeRate'])>0]
    gdea_result={'rows':len(gdea),'date_min':min(x['currentDay'] for x in gdea),'date_max':max(x['currentDay'] for x in gdea),'message_quantity_conflicts':len(conflicts),'conflict_dates':sorted(conflicts),'missing_reported_quantity_rows':len(missing_quantity),'missing_quantity_dates':sorted(missing_quantity),'missing_quantity_filled_with_zero':False,'quantity_field':'exchangeRate (table-defined quantity, not FX)','amount_field':'deal','status_certified':False,'explanation':'Message and reported quantity conflict; individual daily notices still required. Backend user/audit fields are not republished.'}
    archive=next(x for x in meta['sources'] if x['filename'].endswith('.zip'))
    hbea={};duplicates=0
    with zipfile.ZipFile(base/archive['filename']) as z:
        if z.testzip() is not None:
            raise ValueError('Archive CRC check failed')
        files=[i for i in z.infolist() if not i.is_dir()]
        if len(files)!=archive['snapshot_pages']:
            raise ValueError('Archive page count does not match the snapshot manifest')
        for file in files:
            text=z.read(file).decode('utf-8')
            for tr in re.findall(r'<tr\b[^>]*>(.*?)</tr>',text,re.S|re.I):
                cells=[html.unescape(re.sub('<[^>]+>','',x)).strip() for x in re.findall(r'<td\b[^>]*>(.*?)</td>',tr,re.S|re.I)]
                if len(cells)!=9 or cells[0]!='HBEA':continue
                date=dt.date.fromisoformat(cells[1]).isoformat()
                if date in hbea:
                    if hbea[date]!=cells:
                        raise ValueError('Conflicting official rows must not be silently deduplicated')
                    duplicates+=1
                else:hbea[date]=cells
    hbea_result={'rows':len(hbea),'date_min':min(hbea),'date_max':max(hbea),'raw_pages':archive['snapshot_pages'],'identical_duplicate_rows':duplicates,'reported_zero_quantity_rows':sum(reported_decimal(x[6])==0 for x in hbea.values()),'missing_reported_quantity_rows':sum(reported_decimal(x[6]) is None for x in hbea.values()),'missing_filled_with_zero':False,'reported_quantity_column_zero_based':6,'latest_price_is_not_assumed_to_be_close':True,'no_trade_status_certified':False,'all_pages_newly_refetched_this_publication_step':False,'scope':'Same-day original official HTML snapshots freshly reread; no stored audit outcomes used.'}
    return {'status':'freshly_recomputed_from_verified_original_snapshots','cea':cea,'gdea':gdea_result,'hbea':hbea_result,'prediction_skill_tested':False,'old_model_outputs_read':False}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1])
    parser.add_argument('--include-local-exchanges',action='store_true',help='Recheck local exchange originals; these files are deliberately not in GitHub')
    args = parser.parse_args();root=args.root.resolve()
    manifest=json.loads((root/'metadata/raw_manifest.json').read_text(encoding='utf-8'))
    for record in manifest['files']:
        body=(root/record['path']).read_bytes()
        if len(body)!=record['bytes'] or hashlib.sha256(body).hexdigest()!=record['sha256']:
            raise ValueError(f"Checksum or length mismatch: {record['path']}")
    result={'snapshot_date':'2026-10-07','executed_utc':dt.datetime.now(dt.timezone.utc).isoformat(),'analysis_mode':'fresh_parse_of_original_bytes_no_stored_audit_or_training_outputs','legacy_results_status':'UNTRUSTED_NOT_READ','model_training_ran':False,'published_datasets':[],'local_exchange_checks':{'status':'not_run_no_local_snapshot_requested'}}
    specs=[('brent_eia_daily','RBRTEd.xls','USD/barrel'),('wti_eia_daily','RWTCd.xls','USD/barrel'),('henryhub_eia_daily','RNGWHHDd.xls','USD/million Btu')]
    for name,file,unit in specs:
        rel='data/raw/eia/'+file
        workbook=xlrd.open_workbook(root/rel);sheet=workbook.sheet_by_name('Data 1');rows=[]
        for i in range(3,sheet.nrows):
            date,value=sheet.row_values(i)[:2]
            if date=='':continue
            rows.append({'date':xlrd.xldate_as_datetime(date,workbook.datemode).date().isoformat(),'spot_price':value if value!='' else None,'unit':unit})
        record=export_dataset(root,name,rows,rel,{'frequency':'daily','definition':sheet.cell_value(2,1),'missing_price_rows':sum(x['spot_price'] is None for x in rows),'nonpositive_price_rows_retained':sum(x['spot_price'] is not None and x['spot_price']<=0 for x in rows),'not_futures':True,'first_release_vintages_verified':False})
        result['published_datasets'].append(record)
    rel='data/raw/policy_uncertainty/China_Mainland_Paper_EPU.xlsx'
    workbook=openpyxl.load_workbook(root/rel,read_only=True,data_only=True)
    for name,sheets in [('china_mainland_epu',['EPU 1949-1978','EPU 1979-1999','EPU 2000 onwards']),('china_mainland_tpu',['TPU 2000 onwards'])]:
        rows=[]
        for sheet in sheets:
            for cells in workbook[sheet].iter_rows(min_row=2,values_only=True):
                year,month,value=cells[:3]
                if year is None or month is None:continue
                rows.append({'date':dt.date(int(year),int(month),1).isoformat(),'index':value,'source_sheet':sheet})
        result['published_datasets'].append(export_dataset(root,name,rows,rel,{'frequency':'monthly','date_definition':'month identifier, not publication time','same_month_daily_feature_available':False,'first_release_times_verified':False}))
    result['epu_permission_in_original_workbook']=workbook['EPU 2000 onwards']['F2'].value
    result['epu_original_source_attribution']=workbook['EPU 2000 onwards']['F1'].value
    workbook.close()
    rel='data/raw/author_wang_2020/data.xlsx';workbook=openpyxl.load_workbook(root/rel,read_only=True,data_only=True)
    rows=[];first={};second={}
    for date,price in workbook['Sheet1'].iter_rows(values_only=True):
        if date is None:continue
        date=date.date().isoformat();first[date]=Decimal(str(price));rows.append({'date':date,'author_reported_price':price})
    for row in workbook['Sheet2'].iter_rows(values_only=True):
        if row[0] is None:continue
        date=dt.date(int(row[0]),int(row[1]),int(row[2])).isoformat();second[date]=Decimal(str(row[4]))
    workbook.close();common=set(first)&set(second)
    result['author_wang_sheet_check']={'sheet1_observations':len(first),'sheet2_observations':len(second),'common_dates':len(common),'date_price_disagreements':sum(first[d]!=second[d] for d in common),'missing_from_one_sheet':len(set(first)^set(second)),'scope':'within-workbook consistency only; not independent exchange-price certification'}
    result['published_datasets'].append(export_dataset(root,'author_wang_eu_futures',rows,rel,{'frequency':'author-described daily','source_grade':'author-published file with verified platform checksum; market prices not independently authenticated','product_contract_roll_currency_unit_verified':False,'license':'CC BY-NC 3.0'}))
    if args.include_local_exchanges:result['local_exchange_checks']=check_local_exchange_snapshots(root)
    (root/'metadata/audit_results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    (root/'metadata/parsed_manifest.json').write_text(json.dumps({'files':result['published_datasets']},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'published_datasets':[{'dataset':x['dataset'],'rows':x['rows'],'date_min':x['date_min'],'date_max':x['date_max']} for x in result['published_datasets']],'legacy_results_status':result['legacy_results_status'],'author_sheet_check':result['author_wang_sheet_check']},ensure_ascii=False))

if __name__=='__main__':
    main()
