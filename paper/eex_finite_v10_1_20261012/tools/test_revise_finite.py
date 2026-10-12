"""Tests for the finite, plain-paragraph Word edit boundary."""
import unittest
from io import BytesIO
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import revise_finite


def protected(doc):
    return [etree.tostring(x, method='c14n', exclusive=True)
            for x in doc.element.iter() if x.tag in {qn('m:oMath'), qn('w:tbl')}]


class FiniteRevisionTest(unittest.TestCase):
    def test_section_break_is_rejected_before_any_target_is_modified(self):
        doc = Document()
        doc.add_paragraph('Safe deletion.')
        section_p = doc.add_paragraph('Section boundary.')
        section_p._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
        before = etree.tostring(doc.element)
        with self.assertRaises(ValueError):
            revise_finite.revise_document(doc, ['Safe deletion.', 'Section boundary.'], [])
        self.assertEqual(etree.tostring(doc.element), before)

    def test_plain_edit_preserves_protected_content_and_refuses_unsafe_targets(self):
        self.assertTrue(callable(getattr(revise_finite, 'revise_document', None)), 'finite editor is not implemented')
        doc = Document()
        doc.add_paragraph('Repeated explanation.')
        p = doc.add_paragraph()
        p.add_run('No financial ')
        p.add_run('interests.')
        math_p = doc.add_paragraph('Equation: ')
        math = OxmlElement('m:oMath'); run = OxmlElement('m:r'); text = OxmlElement('m:t')
        text.text = 'x'; run.append(text); math.append(run); math_p._p.append(math)
        tab = doc.add_table(rows=1, cols=1); tab.cell(0, 0).text = '-14.51'
        before = protected(doc)
        revise_finite.revise_document(doc, ['Repeated explanation.'], [{'old': 'No financial interests.', 'new': 'No financial or non-financial interests.'}])
        self.assertEqual(protected(doc), before)
        self.assertIn('No financial or non-financial interests.', [p.text for p in doc.paragraphs])
        self.assertNotIn('Repeated explanation.', [p.text for p in doc.paragraphs])
        with self.assertRaises(ValueError):
            revise_finite.revise_document(doc, ['Missing paragraph.'], [])
        with self.assertRaises(ValueError):
            revise_finite.revise_document(doc, ['Equation: x'], [])
        doc.add_paragraph('Duplicate.'); doc.add_paragraph('Duplicate.')
        with self.assertRaises(ValueError):
            revise_finite.revise_document(doc, ['Duplicate.'], [])
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        bad = Document(buffer)
        bad.tables[0].cell(0, 0).text = '999'
        with self.assertRaises(ValueError):
            revise_finite.require_protected_equal(doc, bad)


if __name__ == '__main__':
    unittest.main()
