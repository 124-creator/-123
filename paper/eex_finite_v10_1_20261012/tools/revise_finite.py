"""Finite plain-paragraph edits; preserve all equations, tables and drawings.

This program never imports, executes or changes a market-data estimator.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
import hashlib
import importlib.metadata
import json
from pathlib import Path
import shutil
import zipfile

from docx import Document

from docx.oxml.ns import qn
from lxml import etree

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}


def text_of(element):
    return ''.join(n.text or '' for n in element.iter()
                   if n.tag in {qn('w:t'), qn('m:t')})


def protected_content(doc):
    root = etree.fromstring(etree.tostring(doc.element))
    return {tag: [etree.tostring(e, method='c14n', exclusive=True)
                  for e in root.xpath('.//' + tag, namespaces=NS)]
            for tag in ('m:oMath', 'w:tbl', 'w:drawing')}


def require_protected_equal(before, after):
    if protected_content(before) != protected_content(after):
        raise ValueError('Equation, table or drawing content changed')


def require_plain(paragraph):
    for child in paragraph._p:
        if child.tag == qn('w:pPr'):
            if child.find(qn('w:sectPr')) is not None:
                raise ValueError('Paragraph contains section properties')
            continue
        if child.tag != qn('w:r'):
            raise ValueError('Non-plain paragraph cannot be rewritten')
        if any(e.tag not in {qn('w:rPr'), qn('w:t')} for e in child):
            raise ValueError('Run contains protected non-text content')


def revise_document(doc, deletions, replacements):
    operations = []
    targets = list(deletions) + [r['old'] for r in replacements]
    if len(targets) != len(set(targets)):
        raise ValueError('Repeated edit target')
    # Validate every target before mutating the document.
    for old in targets:
        matches = [p for p in doc.paragraphs if text_of(p._p) == old]
        if len(matches) != 1:
            raise ValueError('Expected one exact paragraph: ' + old[:70])
        require_plain(matches[0])
        operations.append(matches[0])
    before = protected_content(doc)
    for p in operations[:len(deletions)]:
        p._p.getparent().remove(p._p)
    for p, replacement in zip(operations[len(deletions):], replacements):
        first_rpr = next((deepcopy(r._r.rPr) for r in p.runs if r._r.rPr is not None), None)
        for child in list(p._p):
            if child.tag != qn('w:pPr'):
                p._p.remove(child)
        run = p.add_run(replacement['new'])
        if first_rpr is not None:
            run._r.insert(0, first_rpr)
    if protected_content(doc) != before:
        raise ValueError('Protected content changed during paragraph edit')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def replace_once(text, old, new):
    if text.count(old) != 1:
        raise ValueError('Expected one exact Markdown match: ' + old[:70])
    return text.replace(old, new, 1)


def write_new(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('x', encoding='utf-8', newline='') as f:
        f.write(content)


def run(source, output):
    source, output = source.resolve(strict=True), output.resolve(strict=True)
    if source == output:
        raise ValueError('Output must be separate from source')
    plan = json.loads((output / 'FINITE_REVISION_PLAN.json').read_text(encoding='utf-8'))
    for name, expected in plan['source_hashes'].items():
        if sha(source / name) != expected:
            raise ValueError('Frozen source hash mismatch: ' + name)
    interest = plan['interest_replacement']
    main = Document(source / 'MANUSCRIPT_EN_v10.docx')
    old_main = Document(source / 'MANUSCRIPT_EN_v10.docx')
    title = Document(source / 'TITLE_PAGE_v10.docx')
    old_title = Document(source / 'TITLE_PAGE_v10.docx')
    replacements = list(plan['main_replacements']) + [{
        'old': 'Competing interests: ' + interest['old'],
        'new': 'Competing interests: ' + interest['new']}]
    removals = [r['old'] for r in plan['delete_plain_paragraphs']]
    revise_document(main, removals, replacements)
    revise_document(title, [], [interest])
    require_protected_equal(old_main, main)
    require_protected_equal(old_title, title)

    main_md = (source / 'MANUSCRIPT_EN.md').read_text(encoding='utf-8')
    for old in removals:
        main_md = replace_once(main_md, '\n\n' + old + '\n\n', '\n\n')
    for r in replacements:
        main_md = replace_once(main_md, r['old'], r['new'])
    title_md = replace_once((source / 'TITLE_PAGE.md').read_text(encoding='utf-8'),
                            interest['old'], interest['new'])
    # Preflight output collisions before any deliverable writes.
    doc_pairs = [('MANUSCRIPT_EN_v10_1.docx', main), ('TITLE_PAGE_v10_1.docx', title)]
    text_outputs = {'MANUSCRIPT_EN.md': main_md, 'TITLE_PAGE.md': title_md}
    copied = [p for folder in ('figures', 'tables') for p in (source / folder).iterdir() if p.is_file()]
    copied += [source / 'support' / n for n in (
        'PRICE_RESULTS.json', 'PRICE_DESCRIPTIVES.json', 'FULL24_SOURCE_INDEX.json',
        'SOURCE_DEFINITION_CHECK.json')]
    copied += [source / 'SOURCE_CATALOG.md', source / 'ONLINE_RESOURCE_1.md', source / 'ESM_1.pdf']
    destinations = [output / n for n, _ in doc_pairs] + [output / n for n in text_outputs]
    destinations += [output / p.relative_to(source) for p in copied]
    destinations += [output / 'ONLINE_RESOURCE_1_v10_1.docx', output / 'support/FINITE_DOCUMENT_QA.json']
    if any(p.exists() for p in destinations):
        raise FileExistsError('A generated output already exists; refusing overwrite')
    for name, doc in doc_pairs:
        doc.save(output / name)
    for name, content in text_outputs.items():
        write_new(output / name, content)
    for p in copied:
        target = output / p.relative_to(source)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(p, target)
    shutil.copyfile(source / 'ONLINE_RESOURCE_1_v10.docx', output / 'ONLINE_RESOURCE_1_v10_1.docx')
    if sha(source / 'ONLINE_RESOURCE_1_v10.docx') != sha(output / 'ONLINE_RESOURCE_1_v10_1.docx'):
        raise ValueError('Copied supplementary DOCX changed')
    # Verify generated objects by reopening, including the unchanged SI.
    records = []
    for original, revised in [('MANUSCRIPT_EN_v10.docx', 'MANUSCRIPT_EN_v10_1.docx'),
                              ('TITLE_PAGE_v10.docx', 'TITLE_PAGE_v10_1.docx'),
                              ('ONLINE_RESOURCE_1_v10.docx', 'ONLINE_RESOURCE_1_v10_1.docx')]:
        before, after = Document(source / original), Document(output / revised)
        require_protected_equal(before, after)
        with zipfile.ZipFile(output / revised) as z:
            if z.testzip() is not None:
                raise ValueError('Invalid generated DOCX')
        ps = protected_content(after)
        records.append({'file': revised, 'sha256': sha(output / revised),
                        'math_objects': len(ps['m:oMath']), 'table_objects': len(ps['w:tbl']),
                        'drawings': len(ps['w:drawing']), 'protected_xml_exactly_preserved': True})
    for p in copied:
        if sha(p) != sha(output / p.relative_to(source)):
            raise ValueError('Copied artifact changed: ' + str(p))
    before_words = sum(len(text_of(p._p).split()) for p in old_main.paragraphs)
    after_words = sum(len(text_of(p._p).split()) for p in main.paragraphs)
    texts = [text_of(p._p) for p in main.paragraphs]
    abstract = texts[texts.index('Abstract') + 1]
    result = {'basis_commit': plan['basis_commit'], 'documents': records,
              'deleted_redundant_paragraphs': len(removals),
              'deleted_paragraph_words': sum(len(t.split()) for t in removals),
              'main_paragraph_words_before': before_words, 'main_paragraph_words_after': after_words,
              'word_count_scope': 'Whitespace tokens in top-level body paragraphs, excluding table cells',
              'abstract_whitespace_words': len(abstract.split()),
              'frozen_figures_tables_and_aggregate_files_identical': True,
              'supplementary_docx_and_pdf_byte_identical_to_v10': True,
              'market_fits': 0, 'pdf_render_verification': 'Recorded separately after actual native rendering',
              'versions': {n: importlib.metadata.version(n) for n in ('python-docx', 'lxml')}}
    write_new(output / 'support/FINITE_DOCUMENT_QA.json', json.dumps(result, ensure_ascii=False, indent=2))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    run(args.source, args.out)
