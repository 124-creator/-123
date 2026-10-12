"""Deterministic editorial revision of v9. No market code, models or data fitting.
Preserve native equation structures and all previously reported table values.
"""
from __future__ import annotations
import argparse, copy, hashlib, json, re, shutil, zipfile
from pathlib import Path
from collections import Counter
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.opc.constants import RELATIONSHIP_TYPE as RT

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
EXPECTED={'MANUSCRIPT_EN_v9.docx':'ce6b68f22d88454d35a534a99b0b1561d0207a550dc31666f0f37e00c424e789',
'ONLINE_RESOURCE_1_v9.docx':'7e6de1b5e9045d3f0343c323f6e34ffb7209d0dc8fd9d9e64fcff22598a82830'}
TABLE_MAP={3:4,4:6,5:7,6:5}

def xquery(el,path):return etree.fromstring(etree.tostring(el)).xpath(path,namespaces=NS)
def fulltext(el):return ''.join(xquery(el,'.//w:t/text()|.//m:t/text()'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def math_tokens(doc):
 return Counter(''.join(e.itertext()) for e in []) if False else Counter(
   ''.join(e.xpath('.//m:t/text()',namespaces=NS))
   for e in etree.fromstring(etree.tostring(doc.element)).xpath('.//m:oMath',namespaces=NS))
def table_texts(doc):
 return [tuple(fulltext(c._tc) for row in t.rows for c in row.cells) for t in doc.tables]
def replace_text_nodes(el,mapping):
 for t in el.iter(qn('w:t')):
  text=t.text or ''
  for old,new in mapping.items():text=text.replace(old,new)
  t.text=text

def reference_rewrite(el,si=False):
 # References to source-report Table 4 are not references to this manuscript.
 if fulltext(el).startswith('German Emissions Trading Authority'):return
 for t in el.iter(qn('w:t')):
  s=t.text or ''
  s=re.sub(r'\bTable ([3-6])(?![0-9])',lambda m:'Table '+str(TABLE_MAP[int(m.group(1))]),s)
  s=s.replace('Section 3.4','Section 4.3').replace('Section 3.7','Section 4.2')
  if si:
   s=s.replace('Figure 2','Figure S1').replace('Fig. 2','Fig. S1')
   s=s.replace('Figures 2 and 3','Figures S1 and 3')
  t.text=s

def fmt_run(r,size=None):
 r.font.name='Times New Roman'
 for a in ['ascii','hAnsi','cs']:r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+a),'Times New Roman')
 if size is not None:r.font.size=Pt(size)

def rich(p,text):
 # Editable, whole-word runs; never convert long word subscripts into individual letters.
 parts=re.split(r'(\[\[(?:i|b):.*?\]\])',text)
 for part in parts:
  if not part:continue
  m=re.fullmatch(r'\[\[(i|b):(.*?)\]\]',part)
  r=p.add_run(m.group(2) if m else part);fmt_run(r)
  if m:r.italic=m.group(1)=='i';r.bold=m.group(1)=='b'
 return p

def newp(doc,text='',style='Body Text'):
 p=doc.add_paragraph(style=style);rich(p,text);return p

def hyperlink(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');f=OxmlElement('w:rFonts');f.set(qn('w:ascii'),'Times New Roman');f.set(qn('w:hAnsi'),'Times New Roman');pr.append(f)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def setup(doc,header):
 sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
 sec.left_margin=sec.right_margin=Mm(18);sec.top_margin=Mm(20);sec.bottom_margin=Mm(20)
 sec.header_distance=Mm(8);sec.footer_distance=Mm(10)
 for name in ['Normal','Body Text','First Paragraph']:
  if name in doc.styles:
   st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(11.5);st.font.color.rgb=RGBColor(0,0,0)
   pf=st.paragraph_format;pf.line_spacing=1.12;pf.space_after=Pt(6);pf.space_before=Pt(0);pf.first_line_indent=Mm(0);pf.keep_with_next=False;pf.widow_control=True
 for name,size in [('Heading 1',13),('Heading 2',11.5),('Heading 3',11.5)]:
  if name not in doc.styles:doc.styles.add_style(name,1)
  st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
  pf=st.paragraph_format;pf.space_before=Pt(13);pf.space_after=Pt(6);pf.keep_with_next=True;pf.line_spacing=1.0
  if name=='Heading 3':st.font.italic=True
 for name,size in [('Caption',10.5),('Table Note',9.5),('Reference',10.5),('Author',11),('Subtitle',12)]:
  if name not in doc.styles:doc.styles.add_style(name,1)
  st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
  pf=st.paragraph_format;pf.space_after=Pt(5);pf.line_spacing=1.02
  if name=='Caption':st.font.bold=False;pf.keep_with_next=True
  if name=='Reference':pf.left_indent=Mm(4);pf.first_line_indent=Mm(-4)
  if name in ('Author','Subtitle'):pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.styles['Title'].font.name='Times New Roman';doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].font.bold=True
 doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for part in [sec.header,sec.footer]:
  for e in list(part._element):part._element.remove(e)
 hp=sec.header.add_paragraph();hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=hp.add_run(header);fmt_run(r,8)
 fp=sec.footer.add_paragraph();fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=fp.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
 lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-GB')
 rpr=doc.styles['Normal'].element.get_or_add_rPr();rpr.append(lang)
 doc.core_properties.author='Zhongfei Tian; Xiang Wang';doc.core_properties.last_modified_by='Zhongfei Tian'
 doc.core_properties.comments='';doc.core_properties.keywords='EUA auctions; quantity dispersion; reported price deviations'
 # No provenance assertions in file metadata.


def style_math(doc):
 count=0
 for r in doc.element.iter(qn('m:r')):
  t=r.find(qn('m:t'))
  if t is None:continue
  txt=t.text or ''
  if txt in {'M','Z'}:
   pr=r.find(qn('m:rPr'))
   if pr is None:pr=OxmlElement('m:rPr');r.insert(0,pr)
   st=pr.find(qn('m:sty'))
   if st is None:st=OxmlElement('m:sty');pr.append(st)
   st.set(qn('m:val'),'b');count+=1
 # Protect full word labels in original OMML.
 for t in doc.element.iter(qn('m:t')):
  if t.text in {'partial','early','late'}:
   r=t.getparent();pr=r.find(qn('m:rPr'))
   if pr is None:pr=OxmlElement('m:rPr');r.insert(0,pr)
   if pr.find(qn('m:nor')) is None:pr.append(OxmlElement('m:nor'))
 return count

def normalize(doc,si=False):
 # Remove redundant draft bookmarks; copied paragraph bodies and equation trees remain.
 for tag in ['bookmarkStart','bookmarkEnd']:
  for e in list(doc.element.iter(qn('w:'+tag))):e.getparent().remove(e)
 reference_mode=False
 for p in doc.paragraphs:
  txt=p.text.strip()
  if txt in {'References','Additional references'}:reference_mode=True
  elif reference_mode and txt:p.style='Reference'
  math=bool(list(p._p.iter(qn('m:oMath'))));drawing=bool(list(p._p.iter(qn('w:drawing'))))
  if p.style.name not in ['Title','Subtitle','Author','Heading 1','Heading 2','Heading 3','Caption','Table Note','Reference']:
   p.style=doc.styles['Body Text']
  pf=p.paragraph_format
  if not drawing and not re.match(r'^\([S0-9]',txt):
   pf.left_indent=Mm(0);pf.right_indent=Mm(0);pf.first_line_indent=Mm(0)
  if txt.startswith('Note:'):
   p.style='Table Note';pf.keep_with_next=False
  if re.match(r'^(Table (S?\d+)|Fig\. (S?\d+))\b',txt):
   if not re.match(r'^Table S?\d+ (?:reports|supplies|records|and|gives|shows|provides|lists|contains)\b',txt):
    p.style='Caption';pf.keep_with_next=not txt.startswith('Fig.');pf.keep_together=True
    # Consistent caption font; preserve maths contained in original paragraph.
    for r in p.runs:r.bold=False
    prefix=re.match(r'^(?:Table S?\d+|Fig\. S?\d+)',txt)
    if prefix:
     # Most original caption prefix is isolated; formatting per run below.
     n=len(prefix.group(0));pos=0
     for r in p.runs:
      if pos<n:r.bold=True
      pos+=len(r.text)
  if re.match(r'^Panel [A-C]:',txt):
   pf.keep_with_next=True;pf.space_after=Pt(3)
   for r in p.runs:r.bold=True
  if drawing:pf.keep_with_next=True;pf.space_after=Pt(2);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  iseq=math and re.search(r'\((?:S)?[1-5]\)\s*$',txt)
  if iseq:
   pf.keep_together=True;pf.keep_with_next=False;pf.space_before=Pt(7);pf.space_after=Pt(9)
   tabs=p._p.get_or_add_pPr().find(qn('w:tabs'))
   if tabs is not None:p._p.get_or_add_pPr().remove(tabs)
   tabs=OxmlElement('w:tabs')
   for pos,val in [(4900,'center'),(9865,'right')]:
    t=OxmlElement('w:tab');t.set(qn('w:pos'),str(pos));t.set(qn('w:val'),val);tabs.append(t)
   p._p.get_or_add_pPr().append(tabs)
  elif not drawing and not re.match(r'^Panel [A-C]:',txt) and p.style.name in ['Body Text','First Paragraph','Normal']:
   pf.keep_with_next=False;pf.line_spacing=1.12;pf.space_before=Pt(0);pf.space_after=Pt(6)
  size=11.5
  if p.style.name=='Table Note':size=9.5;pf.keep_together=True
  elif p.style.name in ['Caption','Reference']:size=10.5
  elif p.style.name=='Title':size=17
  elif p.style.name=='Subtitle':size=12
  elif p.style.name=='Author':size=11
  elif p.style.name=='Heading 1':size=13
  elif p.style.name in ['Heading 2','Heading 3']:size=11.5
  for r in p.runs:fmt_run(r,size)
  if p.style.name=='Reference':pf.left_indent=Mm(4);pf.first_line_indent=Mm(-4)
 for i,p in enumerate(doc.paragraphs[:-1]):
  if p.style.name in {'Heading 1','Heading 2','Heading 3'}:
   nxt=doc.paragraphs[i+1]
   if len(nxt.text.split())<=150:nxt.paragraph_format.keep_together=True
 # Table widths rescaled to available 174 mm, no screenshot tables.
 for tab in doc.tables:
  tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
  pr=tab._tbl.tblPr
  w=pr.find(qn('w:tblW'))
  if w is not None:w.set(qn('w:type'),'dxa');w.set(qn('w:w'),str(round(174/25.4*1440)))
  grid=list(tab._tbl.tblGrid)
  old=[int(x.get(qn('w:w'),'1')) for x in grid];tot=sum(old) or 1
  widths=[round(v/tot*174/25.4*1440) for v in old]
  for g,ww in zip(grid,widths):g.set(qn('w:w'),str(ww))
  for ri,row in enumerate(tab.rows):
   trpr=row._tr.get_or_add_trPr()
   if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
   if ri==0 and trpr.find(qn('w:tblHeader')) is None:trpr.append(OxmlElement('w:tblHeader'))
   for ci,cell in enumerate(row.cells):
    if ci<len(widths):cell.width=Mm(widths[ci]/1440*25.4)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
     p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.02
     p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=True
     for r in p.runs:fmt_run(r,10);r.bold=ri==0 if r.bold is None else r.bold
  # Header row safely links to the first body row; table cells repeat on continuation.
  for cell in tab.rows[0].cells:
   for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
  # Brief tables stay together; this avoids splitting a result from its label.
  if len(tab.rows)<=(12 if si else 8):
   for row in tab.rows[:-1]:
    for cell in row.cells:
     for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
 # Keep the last table row with an immediately following explanatory note.
 for tab in doc.tables:
  nxt=tab._tbl.getnext()
  if nxt is not None and nxt.tag==qn('w:p') and fulltext(nxt).startswith('Note:'):
   for cell in tab.rows[-1].cells:
    for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
 # Keep a mathematical lead-in with its following display.
 elems=list(doc.element.body)
 for i,e in enumerate(elems[1:],1):
  if e.tag==qn('w:p') and list(e.iter(qn('m:oMath'))) and re.search(r'\((?:S)?[1-5]\)\s*$',fulltext(e)):
   prev=elems[i-1]
   if prev.tag==qn('w:p'):Paragraph(prev,doc).paragraph_format.keep_with_next=True
 # Isolated Greek variable runs retain appropriate scientific typography.
 for p in doc.paragraphs:
  for r in p.runs:
   if re.fullmatch(r'[κθβpdxtY]',r.text.strip() or '-') and p.style.name not in ['Title']:
    r.italic=True
 # Scientific style for scalar symbols in prose and table headers.
 all_paragraphs=list(doc.paragraphs)+[p for tb in doc.tables for row in tb.rows for c in row.cells for p in c.paragraphs]
 token=re.compile(r'[κθβ]|(?<![\w])[dxpGYn](?![\w])')
 for p in all_paragraphs:
  if p.style.name in {'Reference','Title','Heading 1','Heading 2','Heading 3'}:continue
  for r in list(p.runs):
   if not token.search(r.text) or any(c.tag not in {qn('w:rPr'),qn('w:t')} for c in r._r):continue
   text=r.text;pos=0;pieces=[]
   for m in token.finditer(text):
    if m.start()>pos:pieces.append((text[pos:m.start()],False))
    pieces.append((m.group(0),True));pos=m.end()
   if pos<len(text):pieces.append((text[pos:],False))
   for value,is_var in pieces:
    el=copy.deepcopy(r._r)
    for child in list(el):
     if child.tag!=qn('w:rPr'):el.remove(child)
    t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text=value;el.append(t)
    if is_var:
     pr=el.get_or_add_rPr();it=pr.find(qn('w:i'))
     if it is None:it=OxmlElement('w:i');pr.append(it)
     it.set(qn('w:val'),'1')
    r._r.addprevious(el)
   r._r.getparent().remove(r._r)
 style_math(doc)

def caption(doc,text):
 p=newp(doc,text,'Caption')
 # Split prefix for predictable official caption bolding.
 m=re.match(r'^(Table S?\d+(?: \(continued\))?|Fig\. S?\d+)(.*)',text,re.S)
 if m:
  p.clear();r=p.add_run(m.group(1));r.bold=True;fmt_run(r,10.5);r=p.add_run(m.group(2));fmt_run(r,10.5)
 return p

def build_main(source,out,plan):
 doc=Document(source/'MANUSCRIPT_EN_v9.docx');oldmath=math_tokens(doc);oldtables=table_texts(doc)
 blocks=[copy.deepcopy(e) for e in doc.element.body]
 for e in list(doc.element.body):
  if e.tag!=qn('w:sectPr'):doc.element.body.remove(e)
 setup(doc,'QUANTITY DISPERSION AND REPORTED PRICE DEVIATIONS')
 copied=[]
 for op in plan['operations']:
  kind=op['op']
  if kind=='p':
   if op.get('style')=='Caption':caption(doc,op['text'])
   else:newp(doc,op['text'],op.get('style','Body Text'))
  elif kind=='h':newp(doc,op['text'],f'Heading {op["level"]}')
  elif kind=='copy':
   i=op['index'];e=copy.deepcopy(blocks[i]);reference_rewrite(e)
   # Original main headings are not copied; substantive blocks preserve their native objects.
   doc.element.body.insert(len(doc.element.body)-1,e);copied.append(i)
  elif kind=='pagebreak':doc.add_page_break()
  elif kind=='figure':
   p=doc.add_paragraph();p.add_run().add_picture(str(out/'figures'/f'{op["name"]}.png'),width=Mm(174));caption(doc,op['caption'])
  elif kind=='links':
   p=doc.add_paragraph(style='Body Text')
   for i,(label,url) in enumerate(op['items']):
    if i:p.add_run(' | ')
    hyperlink(p,label,url)
  else:raise ValueError(kind)
 normalize(doc)
 if math_tokens(doc)!=oldmath:raise ValueError('A native equation token sequence was lost, duplicated or changed')
 if Counter(table_texts(doc))!=Counter(oldtables):raise ValueError('An existing numerical table changed')
 doc.core_properties.title=plan['title'];doc.save(out/'MANUSCRIPT_EN_v10.docx')
 return {'original_native_math_objects':sum(oldmath.values()),'native_math_token_sequences_preserved':True,'old_physical_tables_preserved':len(oldtables),'copied_blocks':copied}

def rewrite_plain(p,text):
 # Use only on paragraphs without native maths; protect them explicitly.
 if list(p._p.iter(qn('m:oMath'))):raise ValueError('Unsafe text replacement over native maths')
 p.clear();rich(p,text)

def build_si(source,out):
 doc=Document(source/'ONLINE_RESOURCE_1_v9.docx');oldmath=math_tokens(doc);oldtables=table_texts(doc)
 setup(doc,'ONLINE RESOURCE 1 | QUANTITY AND PRICE EVIDENCE')
 for p in doc.paragraphs:
  text=p.text
  reference_rewrite(p._p,si=True)
  if text.startswith('The first path contains the raw-workbook identifiers'):
   rewrite_plain(p,'The input release records workbook identities, full extraction, sample restrictions and cell traces. Analysis releases contain code, protocols, aggregate results and computation receipts. Reproduction starts from the fixed input and eligibility rules, not from filtering to a desired sample count. Provider rights remain applicable; fixed links and the conditions for requesting author-held material are set out in DATA_ACCESS_AND_REPRODUCTION.md and Section S17.')
  if text.startswith('Version:'):rewrite_plain(p,'Version: Final editorial revision, 11 October 2026')
  if text.startswith('Journal:'):rewrite_plain(p,'Journal: Empirical Economics\nAuthors: Zhongfei Tian; Xiang Wang\nAffiliation: Zhengzhou University of Aeronautics, Zhengzhou, Henan, China\nCorresponding author: Zhongfei Tian, 15517837680@163.com\nVersion: Final editorial revision, 11 October 2026')
  if text=='S9 Conditioning-set values accompanying Figure 2':rewrite_plain(p,'S9 Conditioning-set sensitivity and Figure S1')
  if text.startswith('The numerical reimplementations use the same archives'):
   rewrite_plain(p,'Independent numerical implementations use the same archives and are not independent-market replications. Computation counts are distinct from sample sizes. Zhongfei Tian prepared and revised the manuscript; Xiang Wang critically reviewed and checked it and was responsible for final-version approval. The research received no financial support. Generative AI assistance is disclosed in the main methods; the authors retain responsibility for the research and its presentation.')
  if text.startswith('The team reports 1,034'):
   rewrite_plain(p,'The fixed price analysis comprised 1,034 estimation requests: eight baseline/extension fits, 999 restricted bootstrap fits, 21 quarter deletions and six year deletions. Four independent implementation checks of the same extended models agreed in the focal coefficient, CR1 standard error and partial R². Twenty-six artificial tests passed in each recorded Python mode, and the 55 frozen input objects were unchanged. These checks assess implementation and record integrity rather than provide another market sample. Public arithmetic checks of the aggregate outputs and stored test statistics are separately documented.')
  if text.startswith('The targeted local source-definition audit covers'):
   rewrite_plain(p,'The targeted source-definition audit covers the 24 held quarterly reports. File identity and native text were checked, the relevant Table 4 explanatory text was located, and the earliest and latest complete methods paragraphs were additionally examined. The descriptions identify a comparison of auction clearing prices with a secondary-market reference but did not establish an explicit subtraction-and-percentage-denominator equation in the examined text. We retain the official signed number without reconstructing a denominator or interpreting a lower value as a larger auction discount. Table S14 supplies page locators and links to the recorded source URLs; a locator is not evidence of native-event-time authentication.')
  if text.startswith('Note: Page numbers reproduce the targeted local definition'):
   rewrite_plain(p,'Note: Page numbers are those recorded in the targeted source-definition audit. Quarter identifiers link to the recorded official source URLs. FULL24_SOURCE_INDEX.json and the accompanying SOURCE_CATALOG.md retain the file hashes and exact edition links. Three quarters contribute no eligible reference observations; one source-date conflict is excluded in 2025Q1. The original reports and row-level price panels are not republished here. These locators do not authenticate a native tick stream, quote age or complete UTC deadline history.')
  if text.startswith('Data-access statement:'):
   rewrite_plain(p,'Data-access statement: analysis code and aggregate outputs are available through the fixed repository links in DATA_ACCESS_AND_REPRODUCTION.md. The authors hold the matched German price inputs, whose unrestricted redistribution has not been authorized. Access enquiries may be sent to Zhongfei Tian (15517837680@163.com); any provision of these files remains conditional on applicable provider permission and reuse conditions. Official editions can also be sought through the source URLs linked in Table S14. A different downloaded file must not silently replace an identified snapshot. The exact 2026 quantity snapshot is not publicly deposited, and the recorded hash is not itself a data archive.')
  if text.startswith('Note: These are point-estimate influence checks'):
   # No numerical change; author voice is already appropriate.
   pass
  if text.startswith('The earlier display ledgers remain'):
   rewrite_plain(p,'Versioned numerical ledgers retain the mapping from displayed quantities to the fixed results. This revision provides tables/PRICE_DISPLAY_LEDGER.csv and support/AGGREGATE_AUDIT.json, with source-release links in DATA_ACCESS_AND_REPRODUCTION.md. The aggregate audit checks stored summaries and test statistics; it does not replace the row-level analysis or the original source audit.')
  if text.startswith('The table and Figure 2 use'):
   rewrite_plain(p,'Table S6 and Figure S1 use the same fixed values. The full-control row repeats the pooled quantity result for orientation and is not another empirical sample. Main Figure 1 shows interquartile ranges, main Figure 2 shows price-contrast CR1 intervals, and main Figure 3 shows secondary quantity-correlation differences. Their uncertainty and dispersion objects are not interchangeable.')
 # Heading style IDs in v9 represented section headings as Heading 3; normalize semantic level.
 for p in doc.paragraphs:
  if re.match(r'^S\d+\s',p.text):p.style='Heading 2'
  elif p.text=='Additional references':p.style='Heading 2'
 # Insert control figure before S10 (no new values).
 anchor=next(p for p in doc.paragraphs if p.text.startswith('S10 '))
 picp=doc.add_paragraph();picp.add_run().add_picture(str(out/'figures/FigS1.png'),width=Mm(174));anchor._p.addprevious(picp._p)
 cp=caption(doc,'Fig. S1 Residual quantity correlations across conditioning sets. A and C denote log-one-plus and natural-log ratios. Bars are exploratory three-month-block 95% intervals, with 499 draws for reduced controls and 999 for full controls. FE denotes reporting-series and year–month fixed effects. This figure reuses Table S6 and introduces no new estimate')
 anchor._p.addprevious(cp._p)
 # Original titles too long in inherited Heading 2: use modest title and regular metadata.
 for p in doc.paragraphs:
  if p.text=='Online Resource 1':p.style='Title'
  elif p.text.startswith('Quantity dispersion and reported price deviations'):p.style='Subtitle'
 # Hyperlink source-quarter cells. Visible numbers/data remain exactly unchanged.
 index=json.loads((source/'support/FULL24_SOURCE_INDEX.json').read_text())
 sources={r['quarter_id'].replace('_',''):r['source_url'] for r in index['reports']}
 source_table=doc.tables[-1]
 for row in source_table.rows[1:]:
  cell=row.cells[0];key=cell.text.strip()
  if key in sources:
   cell.paragraphs[0].clear();hyperlink(cell.paragraphs[0],key,sources[key])
 normalize(doc,si=True)
 # Clear inherited forced pages and empty draft spacers; retain live figures/equations.
 for p in list(doc.paragraphs):
  p.paragraph_format.page_break_before=False
  for br in list(p._p.iter(qn('w:br'))):
   if br.get(qn('w:type'))=='page':br.getparent().remove(br)
  if not fulltext(p._p).strip() and not list(p._p.iter(qn('w:drawing'))) and not list(p._p.iter(qn('m:oMath'))):
   p._p.getparent().remove(p._p)
 # Keep both S11 panels together by linking every intervening block from its title through the note.
 elems=list(doc.element.body)
 start=next(i for i,e in enumerate(elems) if fulltext(e).startswith('Table S11 '))
 end=next(i for i,e in enumerate(elems[start:],start) if fulltext(e).startswith('Note: All values except'))
 for e in elems[start:end]:
  if e.tag==qn('w:p'):Paragraph(e,doc).paragraph_format.keep_with_next=True
  elif e.tag==qn('w:tbl'):
   for row in Table(e,doc).rows:
    for cell in row.cells:
     for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
 # Keep 24-row source table as one unit if it fits a page; avoids an unlabelled continuation.
 for row in source_table.rows[:-1]:
  for cell in row.cells:
   for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
 if 'Authors: Zhongfei Tian; Xiang Wang' not in fulltext(doc.element) or '15517837680@163.com' not in fulltext(doc.element):raise ValueError('Missing required SI author metadata')
 if math_tokens(doc)!=oldmath:raise ValueError('Supplement mathematics changed')
 if table_texts(doc)!=oldtables:raise ValueError('Supplement table values changed')
 doc.core_properties.title='Online Resource 1: Quantity dispersion and reported price deviations'
 doc.save(out/'ONLINE_RESOURCE_1_v10.docx')
 return {'original_native_math_objects':sum(oldmath.values()),'native_math_token_sequences_preserved':True,'existing_physical_tables_preserved':len(oldtables)}

def build_title(out,title):
 doc=Document();setup(doc,'EMPIRICAL ECONOMICS | TITLE PAGE')
 doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
 for stname in ['Title','Subtitle']:
  st=doc.styles[stname];st.font.italic=False
  for e in list(st.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
  rpr=st.element.find(qn('w:rPr'))
  if rpr is not None:
   for e in list(rpr):
    if e.tag==qn('w:spacing'):rpr.remove(e)
 newp(doc,title,'Title');newp(doc,'Conditional associations and historical scope','Subtitle')
 newp(doc,'Zhongfei Tian · Xiang Wang','Author');newp(doc,'Zhengzhou University of Aeronautics, Zhengzhou, Henan, China','Author')
 newp(doc,'Corresponding author: Zhongfei Tian','Author');newp(doc,'E-mail: 15517837680@163.com','Author')
 for head,text in [
 ('Author contributions','Zhongfei Tian prepared the manuscript and undertook its final revision. Xiang Wang was responsible for critical review, checking and approval of the final version.'),
 ('Funding','The authors received no financial support for this research.'),
 ('Competing interests','The authors declare that they have no financial interests relevant to the content of this article.'),
 ('Data and code','Public code, aggregate outputs, fixed versions and source links are listed in DATA_ACCESS_AND_REPRODUCTION.md. Requests concerning author-held inputs should be addressed to the corresponding author; transfer is conditional on applicable provider permissions and use conditions.')]:
  newp(doc,head,'Heading 1');newp(doc,text)
 normalize(doc);doc.core_properties.title=title;doc.save(out/'TITLE_PAGE_v10.docx')


def main():
 a=argparse.ArgumentParser();a.add_argument('--source',type=Path,required=True);a.add_argument('--out',type=Path,required=True);args=a.parse_args()
 source=args.source;out=args.out;out.mkdir(parents=True,exist_ok=True)
 for name,expected in EXPECTED.items():
  if sha(source/name)!=expected:raise ValueError('Unexpected source bytes '+name)
 plan=json.loads((out/'REVISION_PLAN.json').read_text())
 for f in ['Fig1','Fig3']:
  for ext in ['eps','png','svg']:shutil.copyfile(source/'figures'/f'{f}.{ext}',out/'figures'/f'{f}.{ext}')
 for ext in ['eps','png','svg']:shutil.copyfile(source/'figures'/f'Fig2.{ext}',out/'figures'/f'FigS1.{ext}')
 (out/'support').mkdir(exist_ok=True);(out/'tables').mkdir(exist_ok=True)
 for name in ['PRICE_RESULTS.json','PRICE_DESCRIPTIVES.json','SOURCE_DEFINITION_CHECK.json','FULL24_SOURCE_INDEX.json','AGGREGATE_AUDIT.json']:
  shutil.copyfile(source/'support'/name,out/'support'/name)
 for name in ['PRICE_DESCRIPTIVES.csv','PRICE_RESULT_TABLE.csv','PRICE_DISPLAY_LEDGER.csv']:
  shutil.copyfile(source/'tables'/name,out/'tables'/name)
 qa={'main':build_main(source,out,plan),'supplement':build_si(source,out),'new_market_fits':0,'new_bootstrap_fits':0,'original_inputs_changed':False}
 build_title(out,plan['title'])
 for name in ['MANUSCRIPT_EN_v10','ONLINE_RESOURCE_1_v10','TITLE_PAGE_v10']:
  p=out/(name+'.docx');d=Document(p)
  body_text=fulltext(d.element.body)
  words=re.findall(r"\b[\w]+(?:[’'–-][\w]+)*\b",body_text)
  qa[name]={'sha256':sha(p),'physical_tables':len(d.tables),'embedded_figures':len(d.inline_shapes),'all_text_tokens':len(words),'native_math_objects':sum(math_tokens(d).values())}
 (out/'support/STRUCTURAL_QA.json').write_text(json.dumps(qa,indent=2))
 print(json.dumps(qa,indent=2))
if __name__=='__main__':main()
