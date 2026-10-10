"""Build editable DOCX files from reviewed Markdown; never fit a statistical model."""
from pathlib import Path
import json
import re
import copy
import subprocess
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / '_build'
WORK.mkdir(exist_ok=True)

def configure_reference():
    default = WORK / 'default_reference.docx'
    default.write_bytes(subprocess.check_output(['pandoc', '--print-default-data-file', 'reference.docx']))
    doc = Document(default)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(22)
    sec.left_margin = sec.right_margin = Mm(23)
    sec.header_distance = sec.footer_distance = Mm(10)
    for st in doc.styles:
        if st.type in (1, 2):
            st.font.name = 'Times New Roman'
            st.font.color.rgb = RGBColor(0, 0, 0)
    for nm in ('Normal', 'Body Text', 'First Paragraph'):
        if nm in doc.styles:
            st = doc.styles[nm]
            st.font.size = Pt(11.5)
            pf = st.paragraph_format
            pf.line_spacing = 1.15
            pf.space_before, pf.space_after = Pt(0), Pt(6)
            pf.widow_control = True
    for nm, size in [('Title', 18), ('Subtitle', 11.5), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)]:
        st = next(st for st in doc.styles if st.name == nm)
        st.font.size = Pt(size)
        st.font.bold = nm != 'Subtitle'
        st.paragraph_format.space_before = Pt(10 if nm.startswith('Heading') else 0)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    for nm in ('Title', 'Subtitle'):
        doc.styles[nm].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for nm in ('Caption', 'Image Caption'):
        if nm in doc.styles:
            doc.styles[nm].font.size = Pt(10)
    if 'Hyperlink' in doc.styles:
        doc.styles['Hyperlink'].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles['Hyperlink'].font.underline = False
    path = WORK / 'reference.docx'
    doc.save(path)
    return path


def create_one(source, destination, supplementary=False):
    reference = WORK / 'reference.docx'
    src_text = (ROOT / source).read_text()
    src_text = re.sub(r'\\qquad \(S?\d+\)', '', src_text)
    temp_source = WORK / source
    temp_source.write_text(src_text)
    proc = subprocess.run([
        'pandoc', str(temp_source), '--from=markdown+tex_math_dollars',
        '--resource-path=' + str(ROOT), '--reference-doc=' + str(reference),
        '-o', str(ROOT / destination),
    ], capture_output=True, text=True)
    (WORK / (destination + '.log')).write_text(proc.stderr)
    if proc.returncode:
        raise RuntimeError(proc.stderr)
    doc = Document(ROOT / destination)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = hp.add_run('EUA AUCTION QUANTITY DISPERSION | ' + ('ONLINE RESOURCE 1' if supplementary else 'AUTHOR-REVIEW V6'))
    rr.font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    fp._p.append(field)
    refs = False
    for i, p in enumerate(doc.paragraphs):
        t = p.text
        p.paragraph_format.widow_control = True
        if i == 0:
            p.style = doc.styles['Title']
        if t in ('Conditional association and temporal heterogeneity, 2020–2025',):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
        if t == '1 Introduction' and not supplementary:
            p.paragraph_format.page_break_before = True
        if t in ('Zhongfei Tian · Xiang Wang', 'Zhengzhou University of Aeronautics, Zhengzhou, Henan, China') or t.startswith('Corresponding author:'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(5)
        if t == 'References':
            refs = True
            p.paragraph_format.keep_with_next = True
        elif refs:
            p.paragraph_format.line_spacing = 1.03
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Mm(5)
            p.paragraph_format.first_line_indent = Mm(-5)
            for r in p.runs:
                r.font.size = Pt(10.5)
        if t.startswith(('Table ', 'Panel A:', 'Panel B:', 'Panel C:')):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.size = Pt(10.5)
        if t.startswith('Note:'):
            p.paragraph_format.line_spacing = 1.04
            p.paragraph_format.space_after = Pt(7)
            for r in p.runs:
                r.font.size = Pt(9.5)
        if t.startswith('Fig. '):
            p.paragraph_format.line_spacing = 1.04
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_together = True
            for r in p.runs:
                r.font.size = Pt(9.8)
        if p._p.xpath('.//w:drawing'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            if r.style and r.style.name in ('Verbatim Char', 'Source Code'):
                r.font.name = 'DejaVu Sans Mono'
                r.font.size = Pt(8)
    for ti, table in enumerate(doc.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        heads = [c.text for c in table.rows[0].cells]
        nc = len(table.columns)
        if nc == 7:
            ws = [44,20,20,20,20,20,20]
        elif nc == 5 and heads[0] == 'Year':
            ws = [44,30,30,30,30]
        elif nc == 5 and heads[0] == 'Transformation':
            ws = [34,18,44,18,50]
        elif nc == 5 and 'Period' in heads:
            ws = [18,29,15,20,82]
        elif nc == 5 and any('Block' in h for h in heads):
            ws = [15,23,15,54,57]
        elif nc == 5:
            ws = [18,22,18,25,81]
        elif nc == 3 and heads[0] == 'Symbol':
            ws = [18,70,76]
        elif nc == 3 and 'Condition' in heads[0]:
            ws = [62,51,51]
        elif nc == 3 and 'Deletion' in heads[0]:
            ws = [78,43,43]
        elif nc == 3:
            ws = [44,38,82]
        else:
            ws = [164/nc]*nc
        if len(ws) != len(table.columns):
            raise ValueError(f'Table {ti} width mismatch')
        for j, w in enumerate(ws):
            table.columns[j].width = Mm(w)
        for ri, row in enumerate(table.rows):
            trp = row._tr.get_or_add_trPr()
            trp.append(OxmlElement('w:cantSplit'))
            if ri == 0:
                trp.append(OxmlElement('w:tblHeader'))
            for j, cell in enumerate(row.cells):
                cell.width = Mm(ws[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr()
                mar = OxmlElement('w:tcMar')
                for side, value in [('top', 65), ('bottom', 65), ('left', 70), ('right', 70)]:
                    el = OxmlElement('w:' + side)
                    el.set(qn('w:w'), str(value))
                    el.set(qn('w:type'), 'dxa')
                    mar.append(el)
                cp.append(mar)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.06
                    p.paragraph_format.keep_with_next = (ri < len(table.rows) - 1)
                    if j == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    elif nc == 7 or (nc == 5 and heads[0] == 'Year'):
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    elif nc == 5 and ((heads[0] == 'Transformation' and j in (1,3)) or ('Period' in heads and j in (2,3))):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif nc == 3 and heads[0] == 'Transformation' and j == 1:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(9.5 if nc == 7 else 9.8)
                        if ri == 0:
                            r.bold = True
        tp = table._tbl.tblPr
        for old in tp.findall(qn('w:tblBorders')):
            tp.remove(old)
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'bottom', 'insideH', 'insideV', 'left', 'right'):
            el = OxmlElement('w:' + side)
            el.set(qn('w:val'), 'single' if side in ('top', 'bottom') else 'nil')
            el.set(qn('w:sz'), '6')
            el.set(qn('w:color'), '000000')
            borders.append(el)
        tp.append(borders)
        for cell in table.rows[0].cells:
            cp = cell._tc.get_or_add_tcPr()
            cb = OxmlElement('w:tcBorders')
            edge = OxmlElement('w:bottom')
            edge.set(qn('w:val'), 'single');edge.set(qn('w:sz'), '4')
            cb.append(edge);cp.append(cb)

    # Keep native editable formula objects; place numbers at a stable right tab.
    equation_count = 0
    for p in list(doc.paragraphs):
        blocks = p._p.xpath('./m:oMathPara')
        if not blocks:
            continue
        equation_count += 1
        previous = p._p.getprevious()
        if previous is not None and previous.tag == qn('w:p'):
            Paragraph(previous, p._parent).paragraph_format.keep_with_next = True
        maths = blocks[0].findall(qn('m:oMath'))
        if len(maths) != 1:
            raise ValueError('Expected one display equation per paragraph')
        math = copy.deepcopy(maths[0])
        p._p.remove(blocks[0])
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.tab_stops.clear_all()
        from docx.enum.text import WD_TAB_ALIGNMENT
        p.paragraph_format.tab_stops.add_tab_stop(Mm(82), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(164), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('\t')
        p._p.append(math)
        number = p.add_run('\t(' + ('S' if supplementary else '') + str(equation_count) + ')')
        number.font.name = 'Times New Roman'
        number.font.size = Pt(10.5)
        p.paragraph_format.keep_together = True
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)

    alt_descriptions = [
        'Auction-level medians and interquartile ranges of bid and award reported dispersion ratios in two historical periods. Bars are quartile ranges, not confidence intervals.',
        'Residual correlations for log-one-plus and natural-log ratios across three conditioning sets, with exploratory 95 percent month-block intervals.',
        'Late-minus-early differences in residual correlation for two transformations, with secondary unadjusted exploratory 95 percent intervals.'
    ]
    for i, shape in enumerate(doc.inline_shapes):
        desc = alt_descriptions[i]
        shape._inline.docPr.set('descr', desc)
    props = doc.core_properties
    props.author = 'Zhongfei Tian; Xiang Wang';props.last_modified_by = ''
    props.title = ('Online Resource 1 — ' if supplementary else '') + 'Bid and award dispersion in European carbon auctions'
    props.subject = 'Author-review v6; authors supplied, editorial review, all numerical inputs unchanged'
    props.keywords = 'EUA; auctions; quantity dispersion; conditional association'
    props.comments = 'Author identities supplied; declarations, coauthor approval and data permissions still require confirmation before submission.'
    doc.save(ROOT / destination)
    return {'file': destination, 'tables': len(doc.tables), 'figures': len(doc.inline_shapes),
            'numbered_native_equations': equation_count}

if __name__ == '__main__':
    configure_reference()
    stats = [create_one('MANUSCRIPT_EN.md', 'MANUSCRIPT_EN_v6.docx'),
             create_one('ONLINE_RESOURCE_1.md', 'ONLINE_RESOURCE_1_v6.docx', True),
             create_one('TITLE_PAGE.md', 'TITLE_PAGE_v6.docx')]
    (ROOT / 'support/build_receipt.json').write_text(json.dumps(stats, indent=2))
    print(json.dumps(stats, indent=2))
