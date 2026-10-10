"""Patch reviewed v6.1 DOCX, preserving all pre-existing equations/tables/figures."""
from pathlib import Path
import argparse,copy,hashlib,json,re
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

EXPECTED={'MANUSCRIPT_EN_v6.docx':'c8ab33a6ad282f4239a66c1a0140e0de9b4e0170fb72d99259a848d13d162753','ONLINE_RESOURCE_1_v6.docx':'e0c4da5c1147d5e3a173c75fd79fa094203a6d501adfbc6fbe402e2e366e6670','TITLE_PAGE_v6.docx':'1a72b118ed1a74386ba3cdab62cec2cc64038537287e894fc427bd1a62091192'}

def single(doc,prefix):
    x=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(x)!=1:raise ValueError(f'Expected one paragraph for {prefix}: {len(x)}')
    return x[0]

def replace(doc,prefix,text):
    p=single(doc,prefix)
    if p._p.xpath('.//m:oMath'):raise ValueError('Refuse to flatten existing math')
    fmt=copy.deepcopy(p._p.pPr)
    p.clear();p.add_run(text)
    if fmt is not None:
        old=p._p.pPr
        if old is not None:p._p.remove(old)
        p._p.insert(0,fmt)
    return p

def render_blocks(doc,blocks,anchor=None):
    for item in blocks:
        typ=item['type']
        if typ=='table':
            heads=item['heads'];rows=item['rows'];t=doc.add_table(rows=1,cols=len(heads));t.autofit=False
            for j,h in enumerate(heads):t.rows[0].cells[j].text=h
            for row in rows:
                cells=t.add_row().cells
                for c,v in zip(cells,row):c.text=str(v)
            for j,w in enumerate(item['widths']):t.columns[j].width=Mm(w)
            for i,row in enumerate(t.rows):
                trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
                if i==0:trp.append(OxmlElement('w:tblHeader'))
                for j,c in enumerate(row.cells):
                    c.width=Mm(item['widths'][j]);cp=c._tc.get_or_add_tcPr()
                    mar=OxmlElement('w:tcMar')
                    for side,size in [('top',65),('bottom',65),('left',70),('right',70)]:
                        el=OxmlElement('w:'+side);el.set(qn('w:w'),str(size));el.set(qn('w:type'),'dxa');mar.append(el)
                    cp.append(mar)
                    for p in c.paragraphs:
                        p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
                        p.paragraph_format.line_spacing=1.05
                        p.paragraph_format.keep_with_next=i<len(t.rows)-1
                        for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(9.5);run.bold=i==0
            borders=OxmlElement('w:tblBorders')
            for s in ['top','bottom','left','right','insideH','insideV']:
                e=OxmlElement('w:'+s);e.set(qn('w:val'),'single' if s in ('top','bottom') else 'nil');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');borders.append(e)
            t._tbl.tblPr.append(borders)
            for c in t.rows[0].cells:
                b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');b.append(e);c._tc.get_or_add_tcPr().append(b)
            if anchor is not None:anchor._p.addprevious(t._tbl)
        else:
            style='Heading 3' if typ=='heading' else 'Body Text'
            p=doc.add_paragraph(item['text'],style=next(s for s in doc.styles if s.name==style))
            p.paragraph_format.widow_control=True
            if typ in ('caption','panel'):
                p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.bold=typ=='caption';r.font.size=Pt(10.5)
            elif typ=='note':
                p.paragraph_format.line_spacing=1.04;p.paragraph_format.space_after=Pt(8)
                for r in p.runs:r.font.size=Pt(9.5)
            if anchor is not None:anchor._p.addprevious(p._p)

def fingerprints(doc):
    return {'math':[hashlib.sha256(etree.tostring(e,method='c14n')).hexdigest() for e in doc._element.xpath('.//m:oMath')],
            'tables':[hashlib.sha256(etree.tostring(t._tbl,method='c14n')).hexdigest() for t in doc.tables],
            'inline_images':len(doc.inline_shapes)}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,required=True);ap.add_argument('--out',type=Path,required=True);a=ap.parse_args()
    out=a.out;out.mkdir(parents=True,exist_ok=True);plan=json.loads((out/'support/REVISION_PLAN.json').read_text())
    receipt={'scope':'document integration only; new statistical computations stored separately','documents':{}}
    for name,h in EXPECTED.items():
        path=a.source/name
        if hashlib.sha256(path.read_bytes()).hexdigest()!=h:raise ValueError('Corrected v6.1 source bytes differ '+name)
        doc=Document(path);before=fingerprints(doc)
        if name.startswith('MANUSCRIPT'):
            for prefix,text in plan['replace_main']:replace(doc,prefix,text)
            for anchor,blocks in plan['insert_main_before'].items():render_blocks(doc,blocks,single(doc,anchor))
            for text in plan['new_references']:
                target=single(doc,'European Commission (2024)') if text.startswith('European') else single(doc,'Lovell MC')
                p=doc.add_paragraph(text);p.paragraph_format.line_spacing=1.03;p.paragraph_format.left_indent=Mm(5);p.paragraph_format.first_line_indent=Mm(-5)
                for r in p.runs:r.font.size=Pt(10.5)
                target._p.addprevious(p._p)
            dest='MANUSCRIPT_EN_v7.docx'
        elif name.startswith('ONLINE'):
            replace(doc,'Conditional association and temporal heterogeneity','Historical association and composition sensitivity, with a 2026 time extension')
            replace(doc,'This resource supports the main article', 'This resource supports the integrated manuscript. The original 2020–2025 estimates, descriptions and sensitivity results remain unchanged in Sections S1–S9. Sections S10–S12 add the recorded composition-weighting analysis, separate 2026 re-estimation, post-result finite-resampling diagnostic and external-price data gate. Their scopes and timing are distinguished. A and B remain identical in the original sample; no extra independent result is attributed to B.')
            replace(doc,'The unified numeric ledger is', 'The inherited numeric ledger is tables/INHERITED_NUMERIC_LEDGER.csv; additional values are linked in tables/INTEGRATED_NUMERIC_LEDGER.csv. Historical descriptions and figure inputs are reused, not recalculated. The integration revision preserves the original formulas and tables, adds new-result tables separately and records a document-level structural comparison.')
            render_blocks(doc,plan['append_supplement'])
            render_blocks(doc,[{'type':'heading','text':'Additional references'},*({'type':'p','text':x} for x in plan['new_references'])])
            dest='ONLINE_RESOURCE_1_v7.docx'
        else:
            replace(doc,'Conditional association and temporal heterogeneity','Historical association and composition sensitivity, with a 2026 time extension')
            replace(doc,'Xiang Wang has reviewed the manuscript.','Xiang Wang has reviewed the preceding manuscript. The added supplementary analyses are identified for author review; this file does not imply submission authorization or invent remaining declarations.')
            dest='TITLE_PAGE_v7.docx'
        for p in doc.paragraphs:
            for run in p.runs:
                if 'Author-review v6' in run.text:
                    run.text=run.text.replace('Author-review v6','Author-review v7')
        for sec in doc.sections:
            for p in sec.header.paragraphs:
                for run in p.runs:run.text=run.text.replace('AUTHOR-REVIEW V6','AUTHOR-REVIEW V7')
        doc.core_properties.subject='Integrated author-review version v7; original estimates unchanged'
        doc.core_properties.comments='No funding; Wang reviewed preceding version; new analyses and remaining declarations require author review.'
        doc.save(out/dest);after=fingerprints(doc)
        if before['math']!=after['math']:raise ValueError('Existing mathematical objects changed')
        # Original tables must survive byte-identically and in the same order.
        left=iter(after['tables'])
        for h in before['tables']:
            if not any(x==h for x in left):raise ValueError('Existing table altered or removed')
        if before['inline_images']!=after['inline_images']:raise ValueError('Image count altered')
        receipt['documents'][dest]={'sha256':hashlib.sha256((out/dest).read_bytes()).hexdigest(),'math_objects':len(after['math']),'old_math_objects_identical':True,'old_tables_identical':True,'tables_including_panels':len(after['tables']),'figures':after['inline_images']}
    (out/'support/DOCUMENT_STRUCTURE_QA.json').write_text(json.dumps(receipt,indent=2))
    print(json.dumps(receipt,indent=2))
if __name__=='__main__':main()
