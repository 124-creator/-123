"""Integrate fixed public price evidence into protected DOCX objects; no market fitting.

Inputs must be the three pinned v8.2 documents and the public results/description
folders. Original equations, old result tables and figures are protected. Output
is a new directory; no private row data is read or written by this program.
"""
from __future__ import annotations
import argparse, collections, csv, hashlib, json, re, shutil, zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def objects(doc, query): return [etree.tostring(x,method='c14n') for x in doc._element.xpath(query)]
def table_content(t): return '|'.join(t._tbl.xpath('.//w:t/text()|.//m:t/text()'))
def replace_plain(p,text):
    if p._p.xpath('.//m:oMath|.//w:drawing'): raise ValueError('Refuse replacement of math or drawing')
    p.clear();p.add_run(text)
def replace_run_text(p,old,new):
    # Replacements occur wholly within a text node; never splice mathematical objects.
    found=0
    for t in p._p.xpath('.//w:t'):
        if old in (t.text or ''):t.text=t.text.replace(old,new);found+=1
    if found!=1:raise ValueError('Expected one protected run replacement: '+old)
def style_of(doc,index):return doc.paragraphs[index].style

def new_paragraph(doc,anchor,text,kind,styles):
    p=doc.add_paragraph(style=styles['h3' if kind=='h3' else 'body'])
    anchor.addprevious(p._p)
    if kind=='caption':
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(4)
        p.add_run(text).bold=True
    elif kind=='note':
        p.paragraph_format.space_after=Pt(8)
        run=p.add_run(text);run.font.size=Pt(9)
    else:p.add_run(text)
    if kind=='h3':p.paragraph_format.keep_with_next=True
    return p

def number(v,d=2):return f'{v:.{d}f}'.replace('-','−')
def interval(v,d=2):return '['+number(v[0],d)+', '+number(v[1],d)+']'

def add_table(doc,anchor,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    anchor.addprevious(t._tbl)
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for cells,vals in [(t.rows[0].cells,headers)]+[(t.add_row().cells,row) for row in rows]:
        for i,(c,val) in enumerate(zip(cells,vals)):
            c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.text=str(val)
            p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.line_spacing=1.0
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9.3)
    for c in t.rows[0].cells:
        for r in c.paragraphs[0].runs:r.bold=True
        c.paragraphs[0].paragraph_format.keep_with_next=True
    hdr=OxmlElement('w:tblHeader');hdr.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(hdr)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for k in ['top','bottom','left','right','insideH','insideV']:
        b=OxmlElement('w:'+k);b.set(qn('w:val'),'single' if k in ['top','bottom'] else 'nil')
        if k in ['top','bottom']:b.set(qn('w:sz'),'8')
        borders.append(b)
    pr.append(borders)
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    for c in t.rows[0].cells:
        tcpr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');borders.append(b);tcpr.append(borders)
    return t

def math_run(text,normal=False,bold=False):
    r=OxmlElement('m:r')
    if normal:
        pr=OxmlElement('m:rPr');n=OxmlElement('m:nor');pr.append(n);r.append(pr)
    if bold:
        pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'bi');pr.append(sty);r.append(pr)
    t=OxmlElement('m:t');t.text=text;r.append(t);return r

def math_sub(base,sub,bold=False):
    node=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(math_run(base,bold=bold));s=OxmlElement('m:sub');s.append(math_run(sub,normal=len(sub)>1));node.extend([e,s]);return node

def transpose_lambda():
    node=OxmlElement('m:sSup');e=OxmlElement('m:e');e.append(math_run('λ',bold=True));sp=OxmlElement('m:sup');sp.append(math_run('′'));node.extend([e,sp]);return node

def equation5(doc,anchor,styles):
    p=doc.add_paragraph(style=styles['body']);anchor.addprevious(p._p)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.1),1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.25),2)
    p.add_run('\t')
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(8)
    m=OxmlElement('m:oMath')
    for el in [math_sub('d','t'),math_run(' = a + κ '),math_sub('x','t'),math_run(' + '),transpose_lambda(),math_sub('c','t',bold=True),math_run(' + '),math_sub('η','Y'),math_run(' + '),math_sub('e','t')]:m.append(el)
    p._p.append(m);p.add_run('\t(5)')
    return p

def main_table(doc,anchor,styles,res,desc):
    new_paragraph(doc,anchor,'Table 6 German reported-price context and conditional associations','caption',styles)
    new_paragraph(doc,anchor,'Panel A: Primary-sample descriptive context (240 auctions, 21 quarters; basis points)','p',styles)
    ds=[r for r in desc['table'] if r['sample']=='primary240']
    rows=[[('Signed deviation' if r['outcome']=='signed_bp' else 'Absolute deviation'),number(r['mean']),number(r['median']),interval([r['q25'],r['q75']]),interval([r['minimum'],r['maximum']]),number(r['sample_sd'])] for r in ds]
    add_table(doc,anchor,['Outcome','Mean','Median','Q1–Q3','Range','SD'],rows,[1.45,.65,.65,1.35,1.25,.9])
    new_paragraph(doc,anchor,'Panel B: Fixed price specifications','p',styles)
    names={'primary_signed_240':'Primary: signed','S1_absolute_240':'S1: absolute','S2_no_flags_203':'S2: no row flags','S3_best_bid_236':'S3: best bid only'}
    rows=[]
    for key,label in names.items():
        r=res[key];rows.append([label,f"{r['n']}/{r['quarters']}",number(r['coefficient'])+'\n'+interval(r['CR1_t_Gminus1_CI95']),number(r['IQR_observed_bid_dispersion'],4),number(r['observed_IQR_gap_bps'])+'\n'+interval(r['observed_IQR_gap_CR1_CI95_bps']),number(r['incremental_partial_R2'],4)])
    add_table(doc,anchor,['Specification','n / G','κ [95% CI]','IQR(x)','IQR contrast\n[95% CI], bp','Partial R²'],rows,[1.18,.67,1.47,.60,1.66,.67])
    new_paragraph(doc,anchor,'Note: Panel A gives unconditional summaries of the reported percentage scale, with sample SD using n−1 and linear-interpolation quartiles. The signed observations comprise 142 negative, 92 positive and six printed-zero values. Panel B retains one primary and three secondary models; each baseline and extension uses identical observations. κ is in basis points per unit of log(1 + bid ratio). Intervals use quarter-cluster CR1 covariance and t(G−1), not bootstrap quantiles. Each IQR contrast uses its own observed x range; partial R² is in-sample. The absolute outcome concerns magnitude, not the signed outcome tested in the primary model. Model and inference definitions are in Section 3.7.','note',styles)

def insert_sequence(doc,anchor,seq,styles,res,desc):
    for kind,text in seq:
        if kind=='table':main_table(doc,anchor,styles,res,desc)
        elif kind=='eq':equation5(doc,anchor,styles)
        else:new_paragraph(doc,anchor,text,kind,styles)

def supplementary(doc,anchor,styles,res,desc,sources,accept,plan):
    p=lambda text,kind='p':new_paragraph(doc,anchor,text,kind,styles)
    p('S14 Price estimand, conditioning set and descriptive context','h3')
    p('The response is the official signed numeric percentage-point field multiplied by 100. Printed zero is retained at source precision; it is not a no-trade indicator or certification of identical unrounded prices. The focal covariate is log(1 + reported bid SD / reported bid mean). The fixed baseline contains an intercept, log auction volume, log total and successful bidder counts, log reported cover ratio, the submitted-price maximum minus minimum in EUR, log same-auction clearing price, observed-year indicators, and a last-trade indicator when it varies. The extended model adds only the focal covariate on the same rows. The primary matrix has rank 14 with 226 residual degrees of freedom. The 236-row best-bid model omits its constant reference indicator and has rank 13. The one-sided nature of a best bid, and the fact that reference prices are not midpoints, restrict the interpretation.')
    p('The source file for prices contains total bid volume B, not the reported cover-ratio field R. Eight shared quantities are cross-checked against the original annual EEX files. R is obtained directly from those original workbooks rather than replaced by B/V. No raw SD of awards is added to the price model, and no bid-to-award-to-price mediation relation is estimated.')
    p('Table S11 Reported price-deviation descriptions in the three fixed samples','caption')
    for outcome,label in [('signed_bp','Panel A: Signed deviations'),('absolute_bp','Panel B: Absolute deviations')]:
        p(label)
        rows=[]
        for r in desc['table']:
            if r['outcome']==outcome:
                name={'primary240':'Primary','noflags203':'No row flags','bestbid236':'Best bid only'}[r['sample']]
                rows.append([name,f"{r['n']} / {r['quarters']}",number(r['mean']),number(r['median']),interval([r['q25'],r['q75']]),interval([r['minimum'],r['maximum']]),number(r['sample_sd'])])
        add_table(doc,anchor,['Sample','n / G','Mean','Median','Q1–Q3','Range','SD'],rows,[1.15,.6,.6,.6,1.25,1.3,.75])
    p('Note: All values except n/G are basis points. Equal auction weights are used. Signed negative/positive/printed-zero counts are 142/92/6 in the primary sample, 118/79/6 in the no-row-flag subset, and 138/92/6 in the best-bid subset. These imply printed-zero shares 2.50%, 2.96% and 2.54%. Absolute outcomes preserve the same zeros. Descriptions were computed after the fits on the unchanged private input sets, not selected to improve estimates. The tables do not measure a treatment effect or procurement benefit.','note')
    p('S15 Cluster uncertainty, scale and numerical diagnostics','h3')
    p('The primary uncertainty display uses quarter-cluster CR1 covariance with finite factor G/(G−1) × (n−1)/(n−k), and a t(G−1) reference. The primary signed coefficient alone receives a restricted-null wild-cluster t test with 999 Rademacher sign vectors, SeedSequence([20261011, 2, 240]), and CR1 studentization. The null sets the focal coefficient to zero; no sign vector is redrawn following a failure. All 999 stored draws are valid. Four statistics meet or exceed the absolute observed statistic of 3.5632, giving the two-sided plus-one p-value (4+1)/(999+1) = 0.005. This is a test, not the source of a bootstrap confidence interval.')
    p('The quarter jackknife deletes each of 21 source quarters, recentres its coefficient estimates at their deletion mean, and uses variance (G−1)/G times their summed squared deviations. This is mean-centred CV3J, distinct from the full-estimate-centred CV3. A t(G−1) interval is displayed around the full-sample coefficient. The approaches follow the cluster-jackknife and wild-bootstrap framework discussed by MacKinnon et al. (2023); they remain approximate with few and potentially dependent time clusters. Agreement of these displays does not adjust the whole exploratory history or make the same observations independent replications.')
    p('Table S12 Distinct price scales and inference diagnostics','caption')
    rows=[]
    labels={'primary_signed_240':'Primary: signed','S1_absolute_240':'S1: absolute','S2_no_flags_203':'S2: no row flags','S3_best_bid_236':'S3: best bid only'}
    for key,label in labels.items():
        r=res[key];rows.append([label,number(r['IQR_residualized_bid_dispersion'],4),number(r['fitted_IQR_gap_bps'])+' '+interval(r['IQR_gap_CR1_CI95_bps']),number(r['scaled_condition_number'],2),number(r['max_quarter_share_of_focal_partial_leverage'],4),number(r['focal_partial_leverage_effective_quarters'],2)])
    add_table(doc,anchor,['Specification','x-resid.\nIQR','Residual-IQR contrast\n[95% CI], bp','Scaled\ncondition','Max. quarter\nshare','Info. equiv.\nquarters'],rows,[1.16,.68,1.65,.72,1.24,.8])
    p('Note: Residual-x contrasts differ from the observed-x contrasts in main Table 6. They are not alternatives selected by effect size. Information shares are proportional to within-model residualized x squared, summed by quarter; their reciprocal concentration is a design diagnostic, not a certified effective independent sample size. The primary observed-x contrast is −14.5059 bp; the residual-x contrast is −10.1622 bp. The primary CR1 SE is 22.0923; the CV3J SE is 27.7203 and its coefficient interval [−136.5430, −20.8960]. The single wild-cluster p-value is 0.005. Raw and scaled condition numbers address different numerical scales; no control was dropped to obtain significance.','note')
    p('S16 Source-subset and year-deletion limits','h3')
    p('The four fixed extended specifications and their own identical-row baselines are reported in main Table 6. Only the primary model has the 999-draw wild-cluster test. The two source subsets overlap the primary data; their similar signs are not independent replications. The absolute-outcome coefficient is positive but imprecise, so neither a reduction in magnitude nor a zero effect has been established. No new natural-log price specification or further window search was added after viewing these results.')
    p('Table S13 Fixed leave-one-year-out primary price coefficients','caption')
    rows=[[r['omitted_year'],r['n'],r['quarters'],number(r['coefficient'])] for r in res['primary_signed_240']['deletion_diagnostics']['leave_year_out']]
    add_table(doc,anchor,['Omitted year','Auctions retained','Quarters retained','κ (bp per unit x)'],rows,[1.4,1.5,1.5,1.85])
    p('Note: These are point-estimate influence checks, not six significance tests. Observed-year effects are rebased after deletion; all economic controls are retained. Deleting 2022 reduces the coefficient magnitude from 78.72 to 45.12, a 42.7% comparison of two estimates, not an estimate of the contribution of that year. No claim is made that the deletion-specific intervals exclude zero or that slopes are invariant across years. The 21 quarter deletions and the four models remain available in the fixed RESULTS.json.','note')
    p('The team reports 1,034 original estimation requests (eight baseline/extension fits, 999 restricted bootstrap fits, 21 quarter deletions and six year deletions), followed by four cross-implementation comparisons of the same models. The signed coefficient, CR1 SE and partial R² agreed numerically. Twenty-six artificial tests passed in each recorded Python mode, and 55 source objects were unchanged. These are team computation and verification records, not a claim of a new raw-data refit by the present editorial reviewer. This revision separately checks the public arithmetic, manifest bytes and stored test statistics without access to private event rows.')
    p('S17 Report definitions, locators and availability','h3')
    p('The targeted local source-definition audit covers 24 held quarterly reports: file identity and native text were checked, relevant Table 4 explanatory text was located, and the earliest and latest complete methods paragraphs were additionally read. That audit describes the target as a comparison of auction clearing prices with the secondary-market reference. It did not locate an explicit subtraction-and-percentage-denominator equation in the examined text. This does not prove that no formula occurs elsewhere or in an image. Accordingly, the article retains the official signed number and the phrase “lower official-reported signed deviation”, without reconstructing a numerator, dividing by an inferred reference price, or relabelling the association as a larger discount.')
    p('Table S14 Report-definition page locators and eligibility','caption')
    lookup={r['quarter_id']:r for r in accept['reports']}
    rows=[]
    for r in sources['sources']:
        name=Path(r['report']).stem;q=lookup[name]
        rows.append([name.replace('_',''),r['page_1based'],q['event_rows'],q['strict_reference_date_rows']])
    add_table(doc,anchor,['Source quarter','Definition page','Reported auctions','Eligible auctions'],rows,[1.55,1.5,1.6,1.6])
    p('Note: Page numbers reproduce the targeted local definition record, not a new page-by-page PDF inspection by this reviewer. The machine-readable SOURCE_DEFINITION_CHECK.json and FULL24_SOURCE_INDEX.json carry each exact report URL and SHA-256. Three quarters contribute no eligible reference observations; one additional source-date conflict is excluded in 2025Q1. The underlying official reports and row-level price panels are not republished here. No native tick stream, quote age or complete UTC deadline history is certified by these locators.','note')
    p('Data-access statement: the historical EEX quantity input and aggregate price outputs have fixed repository identifiers. The author-held reported-price rows are not part of the public replication package because bulk redistribution permissions remain unresolved. Reviewers can inspect the released analysis code, sample restrictions and aggregate evidence, but independent refitting requires lawful access to those exact rows. Before submission, the authors must finalize a journal-acceptable access route. The separate 2026 quantity snapshot limitation in Section S11 remains unchanged; no fresh data acquisition was performed for this integration.')

def normalize_docx(path):
    # Deterministic ZIP metadata; no source media or math conversion.
    with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        for name,data in parts.items():
            zi=zipfile.ZipInfo(name,(2026,10,11,12,0,0));zi.compress_type=zipfile.ZIP_DEFLATED;z.writestr(zi,data)

def build(source,public_root,out,plan_path):
    plan=json.loads(plan_path.read_text());out.mkdir(parents=True,exist_ok=True)
    for name,expected in plan['source_hashes'].items():
        if sha(source/name)!=expected:raise ValueError('Protected source differs: '+name)
    res=json.loads((public_root/'research/eua_local_price_execution_20261011/results/RESULTS.json').read_text())
    desc=json.loads((public_root/'research/eua_price_description_followup_20261011/DESCRIPTIVES.json').read_text())
    srcdef=json.loads((public_root/'research/eua_price_description_followup_20261011/SOURCE_DEFINITION_CHECK.json').read_text())
    srcindex=json.loads((public_root/'research/eua_official_cutoff_pilot_20261011/FULL24_SOURCE_INDEX.json').read_text())
    main=Document(source/'MANUSCRIPT_EN_v8_2.docx');si=Document(source/'ONLINE_RESOURCE_1_v8_2.docx');title=Document(source/'TITLE_PAGE_v8_2.docx')
    original={tag:{'math':objects(d,'.//m:oMath'),'tables':[table_content(t) for t in d.tables]} for tag,d in [('main',main),('si',si)]}
    para=main.paragraphs[:];sp=si.paragraphs[:]
    styles={'body':para[10].style,'h3':para[17].style}
    sstyles={'body':sp[10].style,'h3':sp[11].style}
    for ix,text in plan['main_replacements'].items():replace_plain(para[int(ix)],text)
    # Allow the longer integrated introduction/context to flow without a near-empty page.
    para[26].paragraph_format.page_break_before=False
    insert_sequence(main,para[67]._p,plan['methods_insertion'],styles,res,desc)
    insert_sequence(main,para[101]._p,plan['results_insertion'],styles,res,desc)
    insert_sequence(main,para[106]._p,plan['discussion_insertion'],styles,res,desc)
    # Add references in alphabetic order, retain the older cited sources.
    for anchor,text in [(para[132],plan['dehst_reference']),(para[136],plan['new_reference'])]:
        p=new_paragraph(main,anchor._p,text,'p',styles);p.style=para[125].style
        if p._p.pPr is not None:p._p.remove(p._p.pPr)
        if para[125]._p.pPr is not None:p._p.insert(0,deepcopy(para[125]._p.pPr))
        for run in p.runs:run.font.size=Pt(10.5)
    for ix in plan['main_remove_indices']:
        elem=para[ix]._p;elem.getparent().remove(elem)
    for ix,text in plan['si_replacements'].items():replace_plain(sp[int(ix)],text)
    for ix,changes in plan['si_run_replacements'].items():
        for old,new in changes.items():replace_run_text(sp[int(ix)],old,new)
    for cell,text in zip(si.tables[-1].rows[-1].cells,plan['si_evidence_row']):
        if cell._tc.xpath('.//m:oMath'):raise ValueError('Evidence row unexpectedly contains math')
        for p in cell.paragraphs[1:]:p._p.getparent().remove(p._p)
        replace_plain(cell.paragraphs[0],text)
        for r in cell.paragraphs[0].runs:r.font.size=Pt(9)
    supplementary(si,sp[80]._p,sstyles,res,desc,srcdef,srcindex,plan)
    new_paragraph(si,sp[82]._p,plan['dehst_reference'],'p',sstyles)
    new_paragraph(si,si._element.body.sectPr,plan['new_reference'],'p',sstyles)
    # SI retains the original Li reference; remove obsolete former EC2023 gate reference.
    sp[81]._p.getparent().remove(sp[81]._p)
    tp=title.paragraphs[:];replace_plain(tp[0],plan['title']);replace_plain(tp[1],plan['subtitle'])
    replace_plain(tp[13],"Xiang Wang has reviewed the preceding manuscript. This integrated v9 incorporates the completed team price analysis and descriptive context without new market estimation in the editorial review. No funding is confirmed. Actual contributions, competing interests and a lawful data-access arrangement must be completed before submission; preceding review is not represented as approval of all later changes.")
    for d in (main,si,title):
        for section in d.sections:
            for p in section.header.paragraphs:
                for t in p._p.xpath('.//w:t'):
                    if t.text:t.text=t.text.replace('V8.2','V9').replace('V8.1','V9').replace('V8','V9')
        d.core_properties.title=plan['title'];d.core_properties.subject='Integrated author-review v9; completed price evidence, no new fits'
        d.core_properties.author='Zhongfei Tian; Xiang Wang'
        d.core_properties.modified=datetime(2026,10,11,12,0,0,tzinfo=timezone.utc)
    names=[('main',main,'MANUSCRIPT_EN_v9.docx'),('si',si,'ONLINE_RESOURCE_1_v9.docx'),('title',title,'TITLE_PAGE_v9.docx')]
    report={'scope':'Document integration and object integrity only; not a market refit','version':'v9','new_market_fits':0,'new_bootstrap_fits':0,'files':{},'source_hashes':plan['source_hashes']}
    for tag,d,name in names:
        target=out/name;d.save(target);normalize_docx(target);back=Document(target)
        if tag in original:
            old=collections.Counter(original[tag]['math']);new=collections.Counter(objects(back,'.//m:oMath'))
            if old-new:raise ValueError('Original mathematical XML changed: '+tag)
            protected=original[tag]['tables'] if tag=='main' else original[tag]['tables'][:-1]
            updated=[table_content(t) for t in back.tables]
            if any(t not in updated for t in protected):raise ValueError('Original numeric table changed: '+tag)
        report['files'][name]={'sha256':sha(target),'bytes':target.stat().st_size,'math_objects':len(objects(back,'.//m:oMath')),'physical_tables':len(back.tables),'drawings':len(back.inline_shapes)}
    (out/'support').mkdir(exist_ok=True)
    (out/'support/DOCUMENT_QA.json').write_text(json.dumps(report,indent=2))
    if (source/'figures').is_dir():shutil.copytree(source/'figures',out/'figures',dirs_exist_ok=True)
    # Copied sources are public aggregates, never original quotation rows or reports.
    for rel,dest in [('research/eua_local_price_execution_20261011/results/RESULTS.json','PRICE_RESULTS.json'),('research/eua_price_description_followup_20261011/DESCRIPTIVES.json','PRICE_DESCRIPTIVES.json'),('research/eua_price_description_followup_20261011/SOURCE_DEFINITION_CHECK.json','SOURCE_DEFINITION_CHECK.json'),('research/eua_official_cutoff_pilot_20261011/FULL24_SOURCE_INDEX.json','FULL24_SOURCE_INDEX.json')]:shutil.copy2(public_root/rel,out/'support'/dest)
    print(json.dumps(report,indent=2))

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--source',type=Path,required=True);ap.add_argument('--public-root',type=Path,required=True);ap.add_argument('--out',type=Path,required=True);ap.add_argument('--plan',type=Path,required=True)
    a=ap.parse_args();build(a.source,a.public_root,a.out,a.plan)
