"""Export readable manuscript sources and audit presentation, without statistical fits."""
from __future__ import annotations
import argparse, csv, hashlib, json, re, shutil, subprocess
from pathlib import Path
from docx import Document


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')


def main():
    a=argparse.ArgumentParser(description=__doc__)
    a.add_argument('--paper', type=Path, required=True)
    a.add_argument('--public-root', type=Path, required=True)
    a.add_argument('--audit', type=Path, required=True)
    args=a.parse_args(); root=args.paper; inp=args.public_root
    root.joinpath('tables').mkdir(exist_ok=True); root.joinpath('support').mkdir(exist_ok=True)
    audit=json.loads(args.audit.read_text())
    if not audit.get('passed') or audit['new_market_fits']!=0:
        raise ValueError('Expected completed public aggregate audit without refitting')
    if args.audit.resolve() != (root/'support/AGGREGATE_AUDIT.json').resolve():
        shutil.copyfile(args.audit, root/'support/AGGREGATE_AUDIT.json')
    public_files={
        'PRICE_RESULT_TABLE.csv':'research/eua_local_price_execution_20261011/PRICE_RESULT_TABLE.csv',
        'PRICE_DESCRIPTIVES.csv':'research/eua_price_description_followup_20261011/DESCRIPTIVES.csv'}
    for dest,src in public_files.items():shutil.copyfile(inp/src,root/'tables'/dest)
    checks={}; docs={}
    expected=[('MANUSCRIPT_EN_v9.docx','MANUSCRIPT_EN.md',52,10,3),
              ('ONLINE_RESOURCE_1_v9.docx','ONLINE_RESOURCE_1.md',58,16,0),
              ('TITLE_PAGE_v9.docx','TITLE_PAGE.md',0,0,0)]
    for name,mdname,math,tables,drawings in expected:
        p=root/name; doc=Document(p); docs[name]=doc
        actual=(len(doc._element.xpath('.//m:oMath')),len(doc.tables),len(doc.inline_shapes))
        if actual!=(math,tables,drawings):raise ValueError('Unexpected document objects: '+name)
        alltext='\n'.join(doc._element.xpath('.//w:t/text()|.//m:t/text()'))
        for author in ['Zhongfei Tian','Xiang Wang','15517837680@163.com']:
            if author not in alltext:raise ValueError('Author metadata missing: '+name)
        out=subprocess.run(['pandoc',str(p),'-f','docx','-t','markdown','--wrap=none'],capture_output=True,text=True,check=True)
        text=out.stdout
        title=doc.paragraphs[0].text.strip()
        if not text.lstrip().startswith('# '+title):text='# '+title+'\n\n'+text
        image_counter=[0]
        def fix_image(match):
            image_counter[0]+=1
            return '!['+match.group(1)+'](figures/Fig'+str(image_counter[0])+'.png)'
        text=re.sub(r'!\[([^\]]*)\]\((?:media|word/media)/[^)]+\)(?:\{[^}]*\})?',fix_image,text)
        if image_counter[0]!=drawings:raise ValueError('Image export mismatch: '+name)
        # Export is readable evidence, not a build source for native formulas.
        (root/mdname).write_text(text,encoding='utf-8')
        checks[name]={'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'native_math_objects':math,
                     'physical_table_objects':tables,'drawings':drawings,'markdown_title_restored':True}
    md=docs['MANUSCRIPT_EN_v9.docx']; sd=docs['ONLINE_RESOURCE_1_v9.docx']
    paras=[p.text for p in md.paragraphs]; ip=paras.index('Abstract'); abstract=paras[ip+1]
    words=len(abstract.split()); keys=next(p for p in paras if p.startswith('Keywords:'))
    keyword_count=len(keys.split(':',1)[1].split(';'))
    maincaptions=[p for p in paras if re.match(r'^Table \d+ ',p)]
    sicaptions=[p.text for p in sd.paragraphs if re.match(r'^Table S\d+ [A-Z]',p.text)]
    if words!=216 or keyword_count!=6:raise ValueError('Abstract or keywords changed')
    if [int(re.match(r'Table (\d+)',x)[1]) for x in maincaptions]!=list(range(1,7)):
        raise ValueError('Main table captions not sequential')
    if [int(re.match(r'Table S(\d+)',x)[1]) for x in sicaptions]!=list(range(1,15)):
        raise ValueError('Supplement table captions not sequential')
    alltext='\n'.join(paras)
    for bad in ['no price regression was run','no price-premium regression was run','price informativeness remains untested']:
        if bad.lower() in alltext.lower():raise ValueError('Obsolete pending-price statement: '+bad)
    if not any('no financial support' in p for p in paras):raise ValueError('Funding statement lost')
    # A numeric source locator is not independent empirical evidence. Values below
    # are copied from fixed public outputs; arithmetic audit supplies the checks.
    sources={
        'PRICE_RESULTS.json':json.loads((root/'support/PRICE_RESULTS.json').read_text()),
        'PRICE_DESCRIPTIVES.json':json.loads((root/'support/PRICE_DESCRIPTIVES.json').read_text())}
    ledger=[]
    def walk(obj,path,source):
        if isinstance(obj,bool):return
        if isinstance(obj,(int,float)):
            ledger.append({'source_file':'support/'+source,'json_path':path,'value':repr(obj),
                           'role':'Frozen public result; not a new fit or raw-data verification'})
        elif isinstance(obj,dict):
            for k,v in obj.items():walk(v,path+'.'+k,source)
        elif isinstance(obj,list):
            for i,v in enumerate(obj):walk(v,f'{path}[{i}]',source)
    for name,obj in sources.items():walk(obj,'$',name)
    with (root/'tables/PRICE_NUMERIC_SOURCE_LEDGER.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=list(ledger[0]));w.writeheader();w.writerows(ledger)
    inventory=[]
    for name,doc in docs.items():
        for i,e in enumerate(doc._element.xpath('.//w:p')):
            text=''.join(e.xpath('.//w:t/text()|.//m:t/text()'))
            nums=re.findall(r'(?<![A-Za-z])[-−+]?\d+(?:[.,]\d+)*',text)
            if nums:inventory.append({'document':name,'paragraph_xml_index':i,'tokens':' | '.join(nums),'context':text})
    with (root/'tables/DISPLAY_NUMERICS_V9.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=list(inventory[0]));w.writeheader();w.writerows(inventory)
    write_json(root/'support/PRESENTATION_QA.json',{
        'scope':'Document structure, source locators and public-result presentation only',
        'new_market_fits':0,'new_resamples':0,'new_private_row_descriptions':0,
        'documents':checks,'abstract_words':words,'keywords':keyword_count,
        'main_numbered_tables':len(maincaptions),'supplement_numbered_tables':len(sicaptions),
        'main_figures':3,'price_numeric_source_entries':len(ledger),'display_inventory_paragraphs':len(inventory),
        'ledger_is_not_independent_statistical_verification':True,
        'aggregate_audit_checks':audit['arithmetic_and_identity_checks'],
        'pending_author_declarations_preserved':True,'submission_ready':False})
    print(json.dumps({'status':'PRESENTATION_CHECKS_PASSED','abstract_words':words,'keywords':keyword_count,
                      'numeric_source_entries':len(ledger),'main_tables':6,'SI_tables':14,'new_fits':0}))

if __name__=='__main__':main()
